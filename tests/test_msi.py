"""Tests for the spatial MSI engine: ingestion, imaging, analytics, annotation."""
import numpy as np
import pytest

from smile_msi import demo, spatial, multivariate, imaging, isotopes, annotate, preprocess
from smile_msi.cubestore import CubeStore
from smile_msi.msi import MSIDataset


@pytest.fixture(scope="module")
def ds():
    d = demo.make_synthetic(width=28, height=22, seed=11)
    d.prime()
    return d


@pytest.fixture(scope="module")
def peaks(ds):
    return [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)]


def test_cache_token_identity_and_mutation():
    """cache_token is the correctness contract the feature-list / region result caches
    rely on: a stable per-instance identity that bumps on every in-place mutation, so a
    cache can never serve a stale — or a *different slide's* — derived bundle. Builds its
    own datasets so it can't disturb the shared module fixture."""
    a = demo.make_synthetic(width=16, height=12, seed=1)
    b = demo.make_synthetic(width=16, height=12, seed=2)

    # distinct instances never collide — even two structurally identical slides
    assert a.cache_token() != b.cache_token()
    c = demo.make_synthetic(width=16, height=12, seed=1)
    assert a.cache_token() != c.cache_token()

    # stable while nothing mutates; prime() is deterministic, not a mutation
    t0 = a.cache_token()
    assert a.cache_token() == t0
    a.prime()
    assert a.cache_token() == t0

    # orientation: a real rotation bumps, a no-op set does not
    a.set_orientation(a.orientation)
    assert a.cache_token() == t0
    a.rotate90()
    assert a.cache_token() != t0

    # a preprocessing change bumps
    t1 = a.cache_token()
    a.set_preprocessing([])
    assert a.cache_token() != t1

    # release() drops all derived state → must bump
    t2 = a.cache_token()
    a.release()
    assert a.cache_token() != t2


def test_cube_mean_spectrum_cache():
    """cube_mean_spectrum (the fast region-spectrum path) recomputes a CSR.T@vec on every
    region selection; it must memoize per (cube, mask) so re-selecting is instant, and
    invalidate when the slide mutates (cube dropped) so it never serves a stale spectrum."""
    ds = demo.make_synthetic(width=16, height=12, seed=5)
    ds.prime()
    mask = np.zeros(ds.n_pixels, bool)
    mask[: ds.n_pixels // 2] = True

    assert ds.cube_mean_spectrum(mask) is None          # no cube built yet
    ds.build_mz_cube()

    a1 = ds.cube_mean_spectrum(mask)
    assert a1 is not None
    assert ds.cube_mean_spectrum(mask) is a1            # re-select → cached object, no recompute

    other = ds.cube_mean_spectrum(~mask)
    assert other is not a1                              # a different mask is a distinct entry
    np.testing.assert_allclose(ds.cube_mean_spectrum(mask)[1], a1[1])

    # a mutation drops the cube and its mean cache → no stale spectrum
    ds.set_preprocessing([])
    assert ds.cube_mean_spectrum(mask) is None
    ds.build_mz_cube()
    fresh = ds.cube_mean_spectrum(mask)
    assert fresh is not None and fresh is not a1        # rebuilt cube → fresh computation


def test_build_cube_bumps_cache_token():
    """Building (or rebuilding) the m/z cube changes which ion-extraction path ion_vector
    takes for non-picked m/z (binned cube approx vs exact stream), so a feature-list bundle
    cached before the cube must NOT be served after. cache_token() has to change. (review fix #1)"""
    ds = demo.make_synthetic(width=16, height=12, seed=8)
    ds.prime()
    t0 = ds.cache_token()
    ds.build_mz_cube()
    assert ds.cache_token() != t0
    t1 = ds.cache_token()
    ds.build_mz_cube(bin_ppm=ds._bin_ppm * 3 + 5)        # rebuild at a different resolution
    assert ds.cache_token() != t1


def test_cube_mean_spectrum_threadsafe():
    """cube_mean_spectrum's memo is hit from the GUI thread (region select) and worker
    threads (Stats A/B) at once; the LRU pop+insert must be locked or it races into
    'dict changed size during iteration'. Hammer it from 8 threads with 40 distinct masks
    (forcing eviction) and assert no thread raised. (review fix #2)"""
    import threading
    ds = demo.make_synthetic(width=20, height=16, seed=9)
    ds.prime(); ds.build_mz_cube()
    masks = [np.random.RandomState(i).rand(ds.n_pixels) > 0.5 for i in range(40)]
    errors = []

    def worker():
        try:
            for m in masks:
                ds.cube_mean_spectrum(m)
        except Exception as e:                            # noqa: BLE001
            errors.append(e)
    threads = [threading.Thread(target=worker) for _ in range(8)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert not errors, errors


def test_reduce_2d_fast_small_roi_stays_umap():
    """deterministic=False on a small ROI (n_samples < 50) with > 50 peaks must NOT crash
    PCA(50) and silently fall back to t-SNE — the components are clamped to the data so it
    stays UMAP. (review fix #3)"""
    pytest.importorskip("umap")
    import warnings
    X = np.random.RandomState(3).rand(30, 80)             # n_samples=30 < 50, n_features=80 > 50
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        coords, used = multivariate._reduce_2d(X, "umap", 0, deterministic=False)
    assert used == "umap" and coords.shape == (30, 2)


def test_bulk_ibd_read_matches_perpixel(tmp_path):
    """to_ram's guarded bulk .ibd read must produce the exact same dense matrix as the
    per-pixel read for a contiguous continuous slide. (the guard falls back per-pixel for
    any layout it can't prove contiguous — covered by the existing in-RAM to_ram tests,
    whose stores have no read_intensity_block)."""
    pytest.importorskip("pyimzml")
    from pyimzml.ImzMLWriter import ImzMLWriter
    mzs = np.linspace(100.0, 900.0, 40)
    rng = np.random.RandomState(0)
    coords = [(x, y, 1) for y in range(1, 5) for x in range(1, 7)]   # 24 px continuous
    intens = np.array([rng.rand(40).astype(np.float32) for _ in coords])
    path = str(tmp_path / "synth.imzML")
    with ImzMLWriter(path, mode="continuous") as w:
        for c, row in zip(coords, intens):
            w.addSpectrum(mzs, row, c)

    ds = MSIDataset.from_imzml(path)
    block = ds.store.read_intensity_block()
    assert block is not None and block.shape == (24, 40)
    np.testing.assert_allclose(block, intens, rtol=1e-5, atol=1e-6)

    ref = MSIDataset.from_imzml(path)                    # independent per-pixel read of the same file
    ref_mat = np.array([ref.store.get(i)[1] for i in range(24)], dtype=np.float32)
    np.testing.assert_allclose(block, ref_mat, rtol=1e-5, atol=1e-6)
    assert ds.to_ram()                                    # uses the bulk read
    np.testing.assert_allclose(ds.store.matrix, ref_mat, rtol=1e-5, atol=1e-6)


def test_bulk_ibd_read_64bit_writable(tmp_path):
    """64-bit (float64) intensities: the bulk read must cast to float32 in bounded chunks
    (no ~3x transient spike) and return a WRITABLE matrix (np.frombuffer alone is read-only).
    Regression for the batch-2 review's LOW memory/aliasing finding."""
    pytest.importorskip("pyimzml")
    from pyimzml.ImzMLWriter import ImzMLWriter
    mzs = np.linspace(100.0, 900.0, 30)
    rng = np.random.RandomState(1)
    coords = [(x, y, 1) for y in range(1, 4) for x in range(1, 6)]   # 15 px
    intens = np.array([rng.rand(30) for _ in coords])                # float64
    path = str(tmp_path / "synth64.imzML")
    with ImzMLWriter(path, mode="continuous", intensity_dtype=np.float64) as w:
        for c, row in zip(coords, intens):
            w.addSpectrum(mzs, row, c)

    ds = MSIDataset.from_imzml(path)
    assert np.dtype(ds.store._portable.intensityPrecision).itemsize == 8   # 64-bit slide
    block = ds.store.read_intensity_block()
    assert block is not None and block.dtype == np.float32 and block.flags.writeable
    np.testing.assert_allclose(block, intens.astype(np.float32), rtol=1e-5, atol=1e-6)
    block[0, 0] = 123.0                                   # writable round-trip
    assert block[0, 0] == 123.0


def test_reduce_2d_memo_caches_embeddings():
    """_reduce_2d memoizes by exact input + knobs, so re-running an embedding with identical
    inputs is instant. Proven via a non-deterministic UMAP returning bit-identical coords on
    the second call (only possible from the cache); a different input misses; copies served."""
    pytest.importorskip("umap")
    import warnings
    X = np.random.RandomState(0).rand(200, 8)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        a, _ = multivariate._reduce_2d(X, "umap", 0, deterministic=False)
        b, _ = multivariate._reduce_2d(X, "umap", 0, deterministic=False)
        c, _ = multivariate._reduce_2d(X * 2, "umap", 0, deterministic=False)
    assert np.array_equal(a, b) and a is not b       # cache hit, served as a copy
    assert not np.array_equal(a, c)                  # different input → miss


def test_feature_list_cache_invalidation_helper():
    """A lipid-DB change must drop every cached annotation bundle (the cache keys the DB by
    id(), which a re-imported DB could collide with on a freed address). (review fix #4)"""
    pytest.importorskip("pyqtgraph")
    from collections import OrderedDict
    from smile_msi.gui.features import FeaturesTabMixin

    class Stub(FeaturesTabMixin):
        def __init__(self):
            pass
    s = Stub()
    s._invalidate_feature_list_cache()                    # no cache yet → must not error
    s._fl_cache = OrderedDict({("k",): "bundle"})
    s._invalidate_feature_list_cache()
    assert len(s._fl_cache) == 0


def test_micro_centroid_bincount_matches_loop():
    """The vectorized micro-cluster centroid (per-column weighted bincount in
    spatial._tree_from_scores) must stay bit-identical to the original per-cluster
    boolean-mean loop, including empty clusters → zero centroid + zero size."""
    rng = np.random.RandomState(1)
    for n, M, k in [(2000, 40, 6), (200000, 200, 10)]:
        scores = rng.rand(n, k)
        micro = rng.randint(0, M, n)
        micro[micro == 3] = 0                          # leave cluster 3 empty
        cent_old = np.zeros((M, k)); size_old = np.zeros(M)
        for m in range(M):
            sel = micro == m
            if sel.any():
                cent_old[m] = scores[sel].mean(0); size_old[m] = int(sel.sum())
        size_new = np.bincount(micro, minlength=M).astype(float)
        cent_new = np.zeros((M, k))
        for c in range(k):
            cent_new[:, c] = np.bincount(micro, weights=scores[:, c], minlength=M)
        nz = size_new > 0; cent_new[nz] /= size_new[nz, None]
        assert np.array_equal(size_old, size_new)
        assert np.array_equal(cent_old, cent_new)


def test_label_recolour_lut_matches_loop():
    """The NaN-safe palette-LUT recolour in segment._render_label_image must produce the
    exact same RGBA as the original per-cluster `img == cl` loop — including NaN off-tissue
    pixels and any stray negative label staying fully transparent."""
    rng = np.random.RandomState(0)
    pal = [(10, 20, 30), (40, 50, 60), (70, 80, 90), (100, 110, 120)]
    img = rng.randint(0, 5, size=(40, 50)).astype(float)
    img[rng.rand(40, 50) < 0.3] = np.nan               # off-tissue
    img[0, 0] = -1                                      # stray negative label → transparent
    h, wd = img.shape

    old = np.zeros((h, wd, 4), np.ubyte)
    k = int(np.nanmax(img)) + 1 if np.isfinite(img).any() else 0
    for cl in range(k):
        r, g, b = pal[cl % len(pal)]
        old[img == cl] = [r, g, b, 255]

    new = np.zeros((h, wd, 4), np.ubyte)
    finite = np.isfinite(img)
    if finite.any():
        k = int(np.nanmax(img)) + 1
        lut = np.zeros((k + 1, 4), np.ubyte)
        for cl in range(k):
            r, g, b = pal[cl % len(pal)]
            lut[cl + 1] = [r, g, b, 255]
        idx = np.zeros((h, wd), np.intp)
        idx[finite] = img[finite].astype(np.intp) + 1
        new = lut[idx]
    assert np.array_equal(old, new)


def test_reduce_2d_deterministic_toggle():
    """deterministic=True keeps UMAP seeded and reproducible (the export contract);
    deterministic=False still returns a valid 2-D embedding (the interactive fast path,
    here also exercising the >50-dim PCA pre-reduce). Skips cleanly without umap-learn."""
    pytest.importorskip("umap")
    import warnings
    X = np.random.RandomState(2).rand(120, 60)        # >50 dims → hits the PCA pre-reduce
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        a, ua = multivariate._reduce_2d(X, "umap", 0, deterministic=True)
        b, ub = multivariate._reduce_2d(X, "umap", 0, deterministic=True)
        c, uc = multivariate._reduce_2d(X, "umap", 0, deterministic=False)
    assert ua == ub == uc == "umap"
    assert np.allclose(a, b)                           # seeded → reproducible
    assert c.shape == (120, 2)                         # fast path → valid embedding


def test_detect_samples_two_pieces():
    """Auto-detect should split a slide with two separated tissue pieces into two
    disjoint sample masks, ignoring the empty background."""
    H, W = 24, 24
    img = np.zeros((H, W))                 # empty background
    img[3:9, 3:9] = 100.0                  # tissue piece 1
    img[15:21, 15:21] = 100.0              # tissue piece 2 (well separated)
    rr, cc = np.mgrid[0:H, 0:W]

    class FakeDS:
        n_pixels = H * W

        def tic_image(self):
            return img

        def _pixel_rows_cols(self):
            return rr.ravel(), cc.ravel()

    masks = spatial.detect_samples(FakeDS())
    assert len(masks) == 2
    assert not (masks[0] & masks[1]).any()             # disjoint pieces
    # each ~6x6 piece is recovered (smoothing dilates it a little); neither is empty
    assert all(30 <= int(m.sum()) <= 90 for m in masks)
    assert int(masks[0].sum()) >= int(masks[1].sum())  # largest-first


# ----- core ---------------------------------------------------------------- #
def test_geometry_and_shared_axis(ds):
    assert ds.n_pixels == 28 * 22
    assert ds.store.shared_axis() is not None          # continuous mode
    lo, hi = ds.mz_range
    assert lo < hi


def test_mz_bounds_ignores_nonfinite():
    """A corrupt/misread spectrum (inf/nan m/z) must not poison the bounds — a garbage
    range used to leak into the dataset fingerprint and mint a duplicate session/sample
    on every load. Both the shared-axis and per-spectrum scan paths drop non-finite m/z."""
    from smile_msi.msi import MemoryStore
    coords = [(1, 1), (2, 1)]
    axis = np.array([200.0, np.inf, 800.0])            # identical → shared-axis path
    store = MemoryStore(coords, [axis, axis], [np.ones(3), np.ones(3)])
    assert store.shared_axis() is not None
    assert store.mz_bounds() == pytest.approx((200.0, 800.0))

    # differing endpoints (nan) defeat shared-axis detection → per-spectrum scan path
    a = np.array([200.0, 800.0])
    b = np.array([np.nan, 250.0, 900.0])
    store2 = MemoryStore([(1, 1), (2, 1)], [a, b], [np.ones(2), np.ones(3)])
    assert store2.shared_axis() is None
    assert store2.mz_bounds() == pytest.approx((200.0, 900.0))


def test_fingerprint_stable_when_mz_range_is_garbage():
    """The dataset fingerprint keys the managed session; a drifting/garbage m/z range must
    not change it, or the same slide spawns a new session + duplicate cohort sample each
    load. Garbage ranges collapse to a fixed token, leaving the slide identity stable."""
    from smile_msi import library

    class FakeDS:
        coordinates = np.array([[1, 1], [2, 1], [3, 1]])
        n_pixels = 3
        width = 3
        height = 1

        def __init__(self, mz_range):
            self.mz_range = mz_range

    good = library.dataset_fingerprint(FakeDS((200.0, 2000.0)))
    g1 = library.dataset_fingerprint(FakeDS((float("inf"), 1e308)))
    g2 = library.dataset_fingerprint(FakeDS((-4.2e254, 1.3e302)))
    assert good and ":mzNA:" not in good
    assert ":mzNA:" in g1 and g1 == g2                  # every garbage range → one stable id
    assert good != g1
    # the stable parts (pixels/grid/coords) still distinguish real slides
    assert library.dataset_fingerprint(FakeDS((200.0, 2000.0))) == good


def test_ion_image_shape_and_fill(ds, peaks):
    img = ds.ion_image(peaks[0], tol_ppm=50)
    assert img.shape == (ds.height, ds.width)
    assert np.isfinite(img).all() and img.max() > 0


def test_composite_image(ds, peaks):
    comp = ds.composite_image(peaks[:3], tol_ppm=50)
    assert comp.shape == (ds.height, ds.width)
    # composite of 3 ions >= any single one of them, pixelwise
    single = ds.ion_image(peaks[0], tol_ppm=50)
    assert comp.sum() >= single.sum() - 1e-6


def test_ratio_image(ds, peaks):
    r = ds.ratio_image(peaks[0], peaks[1], tol_ppm=50)
    assert r.shape == (ds.height, ds.width) and np.isfinite(r).all() and r.min() >= 0


def test_max_spectrum_skyline(ds):
    axm, mean = ds.mean_spectrum()
    axx, mx = ds.max_spectrum()
    assert axx.shape == axm.shape
    assert (mx >= mean - 1e-9).all()              # skyline >= mean at every m/z
    assert len(ds.pick_peaks(snr=3, min_rel_intensity=0.01, projection="max")) >= 1
    half = np.arange(ds.n_pixels) < ds.n_pixels // 2
    _, region_max = ds.max_spectrum(mask=half)    # skyline of a pixel group
    assert region_max.max() > 0


def test_peak_picking_finds_demo_lipids(ds, peaks):
    # the 10 base lipids should all be present among picked peaks
    assert len(peaks) >= 10
    from smile_msi.demo import _demo_peaks
    for mz, _w, _name in _demo_peaks():
        assert np.min(np.abs(np.array(peaks) - mz)) < mz * 30 / 1e6


def test_normalization_factors(ds):
    f = ds.norm_factors("tic")
    assert f.shape == (ds.n_pixels,)
    assert abs(np.mean(f[f > 0]) - 1.0) < 0.2          # scaled to ~unit mean
    tic = ds.tic()                                      # per-pixel TIC (QC view)
    assert tic.shape == (ds.n_pixels,) and tic.max() > 0


def test_feature_matrix_cumsum_correctness(ds, peaks):
    # the vectorized (cumsum) window-sum must equal an explicit per-spectrum sum
    ds.ensure_features(peaks, tol_ppm=50, reduce="sum")
    mat = ds.feature_matrix("none")
    mz0 = peaks[0]
    win = mz0 * 50 / 1e6
    expected = np.zeros(ds.n_pixels)
    for i in range(ds.n_pixels):
        m, inten = ds.get_spectrum(i)
        a = np.searchsorted(m, mz0 - win, "left")
        b = np.searchsorted(m, mz0 + win, "right")
        expected[i] = inten[a:b].sum()
    assert np.allclose(mat[:, 0], expected, rtol=1e-4, atol=1e-3)


def test_ensure_features_subset_slices_without_re_extracting(ds, peaks, monkeypatch):
    # Removing features from a list is a strict subset of the cached peaks: ensure_features
    # must slice the matching columns out of the cached matrix instead of streaming the whole
    # slide again (the re-extract that froze the GUI thread on every removal).
    full = ds.ensure_features(peaks, tol_ppm=20, reduce="sum").copy()
    keep = [m for i, m in enumerate(peaks) if i not in (1, 3, len(peaks) // 2)]

    calls = {"n": 0}
    real = type(ds).build_features

    def spy(self, *a, **k):
        calls["n"] += 1
        return real(self, *a, **k)

    monkeypatch.setattr(type(ds), "build_features", spy)
    sub = ds.ensure_features(keep, tol_ppm=20, reduce="sum")
    assert calls["n"] == 0                                # no re-stream on removal
    idx = [list(peaks).index(k) for k in keep]
    assert np.array_equal(sub, full[:, idx])             # sliced columns are byte-identical
    assert np.allclose(ds.feature_peaks, keep)

    # a genuine change still rebuilds: different tolerance, and adding a peak back (superset)
    ds.ensure_features(keep, tol_ppm=50, reduce="sum")
    assert calls["n"] == 1
    ds.ensure_features(peaks, tol_ppm=50, reduce="sum")
    assert calls["n"] == 2
    ds.ensure_features(peaks, tol_ppm=20, reduce="sum")   # restore for the shared fixture


def test_mz_cube():
    d = demo.make_synthetic(width=20, height=16, seed=7)
    d.prime()
    peaks = [p["mz"] for p in d.pick_peaks(snr=3, min_rel_intensity=0.01)]
    mz = peaks[1]
    ref = d.ion_vector(mz, tol_ppm=50)            # streaming reference
    d.build_mz_cube(bin_ppm=15)
    axis, cube, _ = d._cube
    assert axis[0] <= d.mz_range[0] + 1 and axis[-1] >= d.mz_range[1] - 1   # full coverage
    cub = d.ion_vector(mz, tol_ppm=50)            # now served from the sparse cube
    assert np.corrcoef(ref, cub)[0, 1] > 0.99


def test_ion_cache(ds, peaks):
    v1 = ds.ion_vector(peaks[0] + 0.5, tol_ppm=50)   # off a picked peak -> streamed + cached
    v2 = ds.ion_vector(peaks[0] + 0.5, tol_ppm=50)
    assert np.array_equal(v1, v2) and len(ds._ion_cache) >= 1


def test_cube_mean_spectrum_matches_cube_rows():
    """Fast ROI/region spectrum path: cube_mean_spectrum is the exact masked row-mean
    of the cube, and None before any cube is built."""
    d = demo.make_synthetic(width=20, height=16, seed=3)
    d.prime()
    mask = np.zeros(d.n_pixels, dtype=bool); mask[::4] = True
    assert d.cube_mean_spectrum(mask) is None          # no cube yet
    d.build_mz_cube(bin_ppm=15)
    _, cube, _ = d._cube
    ax, spec = d.cube_mean_spectrum(mask)
    ref = np.asarray(cube[np.flatnonzero(mask)].sum(0)).ravel() / mask.sum()
    assert np.allclose(spec, ref)
    axn, specn = d.cube_mean_spectrum()                 # mask=None -> all pixels
    assert np.allclose(specn, np.asarray(cube.sum(0)).ravel() / d.n_pixels)


def _shared_axis_dataset(n_px=400, n_ch=350, seed=11, scale=1000.0):
    """A continuous-mode (shared-axis) dataset over a plain in-RAM list store, plus a set
    of peak m/z — the fixture for the to_ram vectorization tests."""
    from smile_msi.msi import MemoryStore
    rng = np.random.RandomState(seed)
    axis = np.linspace(150, 900, n_ch)
    side = int(np.ceil(np.sqrt(n_px)))
    coords = np.array([(x, y) for y in range(1, side + 1) for x in range(1, side + 1)])[:n_px]
    ints = []
    for _ in range(n_px):
        v = np.zeros(n_ch)
        idx = rng.randint(0, n_ch, 40)
        v[idx] = np.abs(rng.rand(40)) * scale
        ints.append(v)
    peaks = np.sort(axis[rng.randint(0, n_ch, 25)])
    return MemoryStore(coords, [axis] * n_px, ints), peaks


def test_to_ram_dense_matches_streaming():
    """Loading into RAM (a dense float32 matrix) must give the same mean spectrum,
    per-pixel TIC/RMS/median and feature matrix as the streaming path — the vectorized
    whole-dataset passes have to equal the per-pixel loop."""
    from smile_msi.msi import DenseMemoryStore, MemoryStore

    store, peaks = _shared_axis_dataset()
    coords = store.coordinates

    ref = MSIDataset(MemoryStore(coords, [store.shared_axis()] * len(store), store._ints))
    ref.prime(); ref.build_features(peaks)

    ram = MSIDataset(store)
    assert ram.to_ram() is True
    assert isinstance(ram.store, DenseMemoryStore) and ram.store.in_memory
    assert ram._dense() is not None
    ram.prime(); ram.build_features(peaks)

    assert np.allclose(ref._mean[1], ram._mean[1], rtol=1e-3, atol=1e-2)
    for k in ("tic", "rms", "median"):
        assert np.allclose(ref._pix[k], ram._pix[k], rtol=1e-3, atol=1e-2), k
    assert np.allclose(ref._feat.matrix, ram._feat.matrix, rtol=1e-3, atol=1e-1)
    # cube built from RAM still serves a faithful ion image
    ram.build_mz_cube(bin_ppm=15)
    assert ram._cube is not None
    assert np.corrcoef(ram.ion_vector(peaks[5], tol_ppm=50),
                       ref._feat.matrix[:, 5])[0, 1] > 0.99


def test_to_ram_budget_guard():
    """to_ram declines (leaves the lazy store in place) when the dense matrix would
    exceed the RAM budget, and succeeds when the budget is ample — so a file too big for
    memory degrades gracefully instead of thrashing."""
    store, _ = _shared_axis_dataset(n_px=100, n_ch=120)
    d = MSIDataset(store)
    assert d.to_ram(max_bytes=1) is False                 # 1-byte budget → refuse
    assert getattr(d.store, "matrix", None) is None       # store unchanged (still lazy/list)
    d.prime()                                             # streaming still works
    assert d._mean is not None
    assert d.to_ram(max_bytes=10**12) is True             # ample budget → now dense


def test_to_ram_holds_raw_so_preprocessing_stays_correct():
    """to_ram snapshots RAW spectra; a preprocessing transform applied afterwards is still
    honored (the dense fast-path disables, streaming from RAM applies the transform),
    matching the lazy path."""
    from smile_msi.msi import MemoryStore

    store, _ = _shared_axis_dataset(n_px=200, n_ch=120, seed=3, scale=50.0)
    coords, axis, ints = store.coordinates, store.shared_axis(), store._ints
    T = [lambda mz, inten: (mz, inten * 3.0)]

    ref = MSIDataset(MemoryStore(coords, [axis] * len(ints), ints))
    ref.set_preprocessing(T); ref.prime()

    ram = MSIDataset(MemoryStore(coords, [axis] * len(ints), ints))
    ram.to_ram()
    ram.set_preprocessing(T)                              # transform set AFTER the snapshot
    assert ram._dense() is None                           # transforms disable the dense fast-path
    ram.prime()
    assert np.allclose(ref._mean[1], ram._mean[1], rtol=1e-4, atol=1e-3)


def test_dense_ion_matches_stream_ion():
    """An arbitrary-m/z ion image on an in-RAM slide is served by the vectorized dense
    column-range reduction (_dense_ion) instead of the per-pixel _stream_ion loop that
    used to run on the GUI thread and freeze the app. The two must be numerically
    identical — same window, same reduce — so the speed-up changes no value. Regression
    for the 'not responding' freeze on a larger dataset."""
    store, peaks = _shared_axis_dataset()
    d = MSIDataset(store)
    assert d.to_ram() is True
    assert d._dense() is not None and d._cube is None      # in-RAM: no cube is built
    axis = d.store.shared_axis()
    # m/z that are NOT picked-peak features (so ion_vector can't take the feature-matrix path)
    targets = [float(axis[len(axis) // 3]),
               float(axis[len(axis) // 2]) + 0.013,
               float(peaks[3])]
    for mz in targets:
        for reduce in ("sum", "mean", "max"):
            dense = d._dense_ion(mz, 50.0, reduce)
            stream = d._stream_ion(mz, 50.0, reduce)
            assert np.allclose(dense, stream, rtol=1e-5, atol=1e-4), (mz, reduce)
    # ion_vector routes through the dense path (no cube) and equals the streamed reference
    iv = d.ion_vector(targets[0], tol_ppm=50.0)
    assert np.allclose(iv, d._stream_ion(targets[0], 50.0, "sum") / d.norm_factors("none"))


def test_build_mz_cube_chunked_equivalence():
    """The chunked/bounded-memory cube assembly is numerically identical to a direct
    per-point accumulation — across fine bins (no within-bin collisions), coarse bins
    (heavy collisions, the summation path) and a min_intensity gate, for BOTH continuous
    (shared-axis) and processed (ragged) data. Pins the refactor that replaced the
    list-of-arrays + concatenate + per-pixel bincount with vstacked CSR chunks."""
    from smile_msi.msi import _bin_edges, MemoryStore

    def ref_cube(d, minint):
        caxis, cube, _ = d._cube
        edges = _bin_edges(caxis)
        nb = len(caxis)
        ref = np.zeros((d.n_pixels, nb))
        for i in range(d.n_pixels):
            m, inten = d._read(i)
            bi = np.searchsorted(edges, m, side="right") - 1
            ok = (bi >= 0) & (bi < nb) & (inten > minint)
            np.add.at(ref[i], bi[ok], inten[ok])
        return cube.toarray(), ref

    # continuous / shared-axis: fine bins, collision-forcing coarse bins, and a gate
    store, _ = _shared_axis_dataset(n_px=120, n_ch=200, seed=5)
    dc = MSIDataset(store)
    for bin_ppm, minint in [(15.0, 0.0), (8000.0, 0.0), (15.0, 50.0)]:
        dc._cube = None
        dc.build_mz_cube(bin_ppm=bin_ppm, min_intensity=minint)
        got, ref = ref_cube(dc, minint)
        assert np.allclose(got, ref, rtol=1e-4, atol=1e-3), (bin_ppm, minint)

    # processed / ragged: per-pixel m/z, coarse bins so 200.0 & 200.05 collide in one bin
    coords = np.array([(x + 1, 1) for x in range(6)])
    mzs = [np.array([200.0, 200.05, 400.0, 700.0]) for _ in range(6)]
    ints = [np.array([1.0, 2.0, 5.0, 3.0]) * (k + 1) for k in range(6)]
    dp = MSIDataset.from_arrays(coords, mzs, ints)
    dp.build_mz_cube(bin_ppm=3000.0)
    got, ref = ref_cube(dp, 0.0)
    assert np.allclose(got, ref, rtol=1e-4, atol=1e-3)


def _write_legacy_npz(path, axis, csc, edges, fp, mean=None, pix=None):
    """Write a pre-Zarr ``.cache.npz`` cube cache — the format save_cube no longer writes but
    load_cube still reads and migrates. Used to exercise the npz→Zarr migration path."""
    mat = csc.tocsc()
    arrs = {"axis": np.asarray(axis), "edges": np.asarray(edges),
            "data": mat.data, "indices": mat.indices, "indptr": mat.indptr,
            "shape": np.asarray(mat.shape, dtype=np.int64),
            "fingerprint": np.array(str(fp))}
    if mean is not None:
        arrs["mean"] = np.asarray(mean[1] if isinstance(mean, tuple) else mean)
    if pix is not None:
        for k in ("tic", "rms", "median"):
            if pix.get(k) is not None:
                arrs[k] = np.asarray(pix[k])
    np.savez(path, **arrs)


def test_cube_sidecar_roundtrip(tmp_path):
    """save_cube persists the fast cube + prime stats to the chunked Zarr store, load_cube
    returns it as a lazily-read CubeStore, and a mismatched slide (fingerprint / pixel-count
    guard) is rejected. The legacy npz sidecar is no longer written."""
    import os

    from smile_msi import session
    from smile_msi.cubestore import CubeStore

    d = demo.make_synthetic(width=18, height=14, seed=8)
    d.prime(); d.build_mz_cube(bin_ppm=15)
    spath = str(tmp_path / "s.json")
    fp = "slideFP"
    session.save_cube(spath, d._cube, mean=d._mean, pix=d._pix, fingerprint=fp)
    assert os.path.exists(session.cube_zarr_path(spath))          # zarr store written
    assert not os.path.exists(session.cube_sidecar_path(spath))   # npz no longer written

    got = session.load_cube(spath, fp, d.n_pixels)
    assert got is not None
    store = got["cube"]
    assert isinstance(store, CubeStore)                       # zarr-first → lazy store
    axis0, csc0, edges0 = d._cube
    assert np.allclose(store.axis, axis0) and np.allclose(store.edges, edges0)
    assert store.shape == csc0.shape
    # the lazily-read store reconstructs the cube byte-for-byte
    full = store.column_slice(0, store.nbins)
    assert np.array_equal(full.toarray(), csc0.toarray())
    assert np.allclose(got["mean"][1], d._mean[1])
    assert np.allclose(got["pix"]["tic"], d._pix["tic"])
    store.close()
    # guards: a different slide (fingerprint) or pixel count is ignored, never misapplied
    assert session.load_cube(spath, "other-slide", d.n_pixels) is None
    assert session.load_cube(spath, fp, d.n_pixels + 1) is None


def _cube_probe_mzs(d, k=6):
    """A handful of m/z that are NOT picked-peak features (so ion_vector takes the cube
    path, not the feature matrix) and not in RAM — spread across the cube's axis."""
    axis = d._cube[0] if not isinstance(d._cube, CubeStore) else d._cube.axis
    lo, hi = d.mz_range
    return [float(axis[i]) for i in np.linspace(len(axis) // 10, len(axis) - 2, k).astype(int)]


def test_cube_zarr_bit_identical_and_npz_migration(tmp_path):
    """Acceptance gate: a reopened cube gives numbers IDENTICAL to the freshly-built in-RAM
    CSC. Compares ion_vector (sum/mean/max) and cube_mean/max_spectrum across the in-RAM
    tuple, the Zarr reload, and a legacy npz cache migrated to Zarr — all with np.array_equal
    (bit-for-bit, not just close)."""
    import os

    from smile_msi import session

    d = demo.make_synthetic(width=22, height=18, seed=5)
    d.prime(); d.build_mz_cube(bin_ppm=15)
    assert d._dense() is None                       # not in RAM → ion_vector uses the cube
    tuple_cube = d._cube
    axis0, csc0, edges0 = tuple_cube

    spath = str(tmp_path / "s.json")
    fp = "fp-bitexact"
    session.save_cube(spath, d._cube, mean=d._mean, pix=d._pix, fingerprint=fp)

    mzs = _cube_probe_mzs(d)
    for mz in mzs:
        assert d._feature_index(mz) is None         # confirm: not a feature → cube path

    # reference numbers from the in-RAM tuple
    ref_ion = {(round(mz, 6), r): d.ion_vector(mz, tol_ppm=60, reduce=r)
               for mz in mzs for r in ("sum", "mean", "max")}
    mask = np.zeros(d.n_pixels, dtype=bool); mask[::3] = True
    ref_mean_all = d.cube_mean_spectrum()[1]
    ref_mean_msk = d.cube_mean_spectrum(mask)[1]
    ref_max_all = d.cube_max_spectrum()[1]
    ref_max_msk = d.cube_max_spectrum(mask)[1]

    def _check(label):
        for mz in mzs:
            for r in ("sum", "mean", "max"):
                got = d.ion_vector(mz, tol_ppm=60, reduce=r)
                assert np.array_equal(got, ref_ion[(round(mz, 6), r)]), f"{label} ion {mz} {r}"
        assert np.array_equal(d.cube_mean_spectrum()[1], ref_mean_all), f"{label} mean all"
        assert np.array_equal(d.cube_mean_spectrum(mask)[1], ref_mean_msk), f"{label} mean mask"
        assert np.array_equal(d.cube_max_spectrum()[1], ref_max_all), f"{label} max all"
        assert np.array_equal(d.cube_max_spectrum(mask)[1], ref_max_msk), f"{label} max mask"

    # (1) Zarr reload — the preferred, lazily-read CubeStore
    got = session.load_cube(spath, fp, d.n_pixels)
    assert isinstance(got["cube"], CubeStore)
    d._cube = got["cube"]
    _check("zarr")
    d._close_cube()

    # (2) legacy npz migration — drop the zarr, plant an old-format npz cache, and confirm
    #     load_cube migrates it to a fresh Zarr store (lazy CubeStore) with identical numbers.
    os.remove(session.cube_zarr_path(spath))
    _write_legacy_npz(session.cube_sidecar_path(spath), axis0, csc0, edges0, fp,
                      mean=d._mean, pix=d._pix)
    got = session.load_cube(spath, fp, d.n_pixels)
    assert isinstance(got["cube"], CubeStore)                      # migrated npz → zarr
    assert os.path.exists(session.cube_zarr_path(spath))           # fresh zarr written
    assert not os.path.exists(session.cube_sidecar_path(spath))    # superseded npz removed
    d._cube = got["cube"]
    _check("migrated")
    d._close_cube()

    # restore the in-RAM tuple and confirm it still matches (sanity: refs weren't mutated)
    d._cube = tuple_cube
    _check("tuple")


def test_cube_store_staleness_and_corrupt_safety(tmp_path):
    """CubeStore.open returns None (never a mismatched cube) on a wrong fingerprint / pixel
    count / too-new format, and load_cube degrades a corrupt zarr store to "no cache" rather
    than crashing. A stale npz that mismatches the slide is ignored (not migrated)."""
    import os

    from smile_msi import session
    from smile_msi import cubestore

    d = demo.make_synthetic(width=16, height=12, seed=2)
    d.prime(); d.build_mz_cube(bin_ppm=15)
    spath = str(tmp_path / "s.json")
    fp = "guarded"
    session.save_cube(spath, d._cube, mean=d._mean, pix=d._pix, fingerprint=fp)
    zpath = session.cube_zarr_path(spath)

    # direct guards on the store
    assert cubestore.CubeStore.open(zpath, fingerprint=fp, n_pixels=d.n_pixels) is not None
    assert cubestore.CubeStore.open(zpath, fingerprint="nope", n_pixels=d.n_pixels) is None
    assert cubestore.CubeStore.open(zpath, fingerprint=fp, n_pixels=d.n_pixels + 1) is None
    assert cubestore.CubeStore.open(zpath + ".missing", fingerprint=fp, n_pixels=d.n_pixels) is None

    # a too-new on-disk format is rejected even for the right slide
    import unittest.mock as mock
    with mock.patch.object(cubestore, "FORMAT_VERSION", cubestore.FORMAT_VERSION + 1):
        assert cubestore.CubeStore.open(zpath, fingerprint=fp, n_pixels=d.n_pixels) is None

    # corrupt the zarr store with no npz present → "no cache" (rebuild), never an exception
    with open(zpath, "r+b") as f:
        f.truncate(os.path.getsize(zpath) // 2)
    assert session.load_cube(spath, fp, d.n_pixels) is None
    assert not os.path.exists(session.cube_sidecar_path(spath))     # confirm: no npz is written

    # a legacy npz whose fingerprint mismatches this slide is ignored (never migrated)
    axis0, csc0, edges0 = d._cube
    os.remove(zpath)
    _write_legacy_npz(session.cube_sidecar_path(spath), axis0, csc0, edges0, "OTHER-SLIDE")
    assert session.load_cube(spath, fp, d.n_pixels) is None
    assert os.path.exists(session.cube_sidecar_path(spath))         # stale npz left untouched
    assert not os.path.exists(zpath)                                # nothing migrated


def test_cube_store_atomic_write(tmp_path):
    """Writing is atomic (tmp + os.replace): a stale leftover .tmp doesn't break a fresh
    create, and a successful create leaves no .tmp behind."""
    import os

    d = demo.make_synthetic(width=14, height=10, seed=9)
    d.prime(); d.build_mz_cube(bin_ppm=15)
    axis, csc, edges = d._cube
    zpath = str(tmp_path / "s.cube.zarr")

    # plant a junk leftover from a hypothetical killed prior write
    with open(zpath + ".tmp", "wb") as f:
        f.write(b"half-written garbage")
    store = CubeStore.create(zpath, axis, csc, edges, fingerprint="atomic",
                             n_pixels=d.n_pixels)
    assert store is not None
    assert not os.path.exists(zpath + ".tmp")          # tmp consumed by the atomic replace
    # the freshly created store reads back correctly despite the stale tmp
    full = store.column_slice(0, store.nbins)
    assert np.array_equal(full.toarray(), csc.toarray())
    store.close()


def _streaming_build_dataset(seed=21):
    """A small *sparse* shared-axis dataset (≈40 nonzeros/pixel) — light enough that even a
    pathologically small RAM budget (forcing many column bands) builds fast, and with bin
    collisions so the per-pixel float64→float32 sum is exercised."""
    from smile_msi.msi import MemoryStore
    rng = np.random.RandomState(seed)
    n_px, n_ch = 300, 280
    axis = np.linspace(150.0, 900.0, n_ch)
    side = int(np.ceil(np.sqrt(n_px)))
    coords = np.array([(x, y) for y in range(1, side + 1)
                       for x in range(1, side + 1)])[:n_px]
    ints = []
    for _ in range(n_px):
        v = np.zeros(n_ch)
        idx = rng.randint(0, n_ch, 40)            # repeats → bin collisions within a pixel
        v[idx] = np.abs(rng.rand(40)) * 1000.0
        ints.append(v)
    return MemoryStore(coords, [axis] * n_px, ints)


@pytest.mark.parametrize("max_ram_bytes", [1 << 11, 1 << 16, 1 << 24])
def test_cube_streaming_build_bit_identical(tmp_path, max_ram_bytes):
    """The out-of-core streaming build (build_mz_cube(out_path=...)) must produce a cube
    BYTE-identical to the in-RAM build, for any RAM budget (i.e. any number of column
    bands). Checks the CSC arrays, the stored mean numerator, and ion/mean/max results."""
    # in-RAM reference
    ref = MSIDataset(_streaming_build_dataset()); ref.prime()
    ref.build_mz_cube(bin_ppm=20, min_intensity=5.0)
    axis0, csc0, edges0 = ref._cube

    # streaming build of the same data, straight to a CubeStore on disk
    d = MSIDataset(_streaming_build_dataset()); d.prime()
    zp = str(tmp_path / f"s_{max_ram_bytes}.cube.zarr")
    store = d.build_mz_cube(bin_ppm=20, min_intensity=5.0, out_path=zp,
                            fingerprint="strm", max_ram_bytes=max_ram_bytes)
    assert isinstance(store, CubeStore) and d._cube is store

    full = store.column_slice(0, store.nbins)
    assert np.array_equal(full.data, csc0.data)
    assert np.array_equal(full.indices.astype(csc0.indices.dtype), csc0.indices)
    assert np.array_equal(full.indptr.astype(csc0.indptr.dtype), csc0.indptr)
    assert np.array_equal(store.axis, axis0) and np.array_equal(store.edges, edges0)
    # colsum (the cube_mean_spectrum numerator) is the exact csc.T @ ones
    assert np.array_equal(store.colsum(),
                          np.asarray(csc0.T @ np.ones(csc0.shape[0])).ravel())

    # ion images + spectra identical to the in-RAM path, across several m/z and reduces
    mzs = [float(axis0[j]) for j in np.linspace(len(axis0) // 8, len(axis0) - 2, 5).astype(int)]
    for mz in mzs:
        for red in ("sum", "mean", "max"):
            assert np.array_equal(ref.ion_vector(mz, tol_ppm=80, reduce=red),
                                  d.ion_vector(mz, tol_ppm=80, reduce=red)), (mz, red)
    assert np.array_equal(ref.cube_mean_spectrum()[1], d.cube_mean_spectrum()[1])
    assert np.array_equal(ref.cube_max_spectrum()[1], d.cube_max_spectrum()[1])
    mask = np.zeros(ref.n_pixels, dtype=bool); mask[::4] = True
    assert np.array_equal(ref.cube_mean_spectrum(mask)[1], d.cube_mean_spectrum(mask)[1])

    # prime stats are carried into the streamed store (reopen skips the prime pass)
    mean, pix = store.stored_extras()
    assert mean is not None and pix is not None and np.allclose(pix["tic"], d._pix["tic"])
    store.close()


def test_cube_streaming_build_reopen_and_empty(tmp_path):
    """A streamed cube reopens under its fingerprint guard, and an all-thresholded (empty)
    cube streams to a valid zero-nnz store rather than crashing."""
    d = MSIDataset(_streaming_build_dataset()); d.prime()
    zp = str(tmp_path / "r.cube.zarr")
    store = d.build_mz_cube(bin_ppm=20, min_intensity=0.0, out_path=zp, fingerprint="fp")
    nnz = store.nnz
    store.close()
    # reopens with the matching guard; rejects a mismatch
    reopened = CubeStore.open(zp, fingerprint="fp", n_pixels=d.n_pixels)
    assert reopened is not None and reopened.nnz == nnz
    reopened.close()
    assert CubeStore.open(zp, fingerprint="other", n_pixels=d.n_pixels) is None

    # everything below threshold → empty cube, still a valid store
    e = MSIDataset(_streaming_build_dataset()); e.prime()
    ze = str(tmp_path / "e.cube.zarr")
    es = e.build_mz_cube(bin_ppm=20, min_intensity=1e12, out_path=ze, fingerprint="e")
    assert es.nnz == 0
    assert np.array_equal(es.ion(float(es.axis[5]), 80.0, "sum", e.n_pixels),
                          np.zeros(e.n_pixels))
    assert np.array_equal(es.colsum(), np.zeros(es.nbins))
    es.close()


def test_parallel_streaming_matches_serial():
    """The threaded streaming passes (prime/build_features/build_mz_cube) must produce
    results identical to the serial path. The default MemoryStore is serial, so opt a
    subclass into parallel reads (RAM is thread-safe) and compare both datasets."""
    from smile_msi.msi import MemoryStore

    rng = np.random.RandomState(5)
    n_px, n_ch = 1500, 600          # >1024 so the parallel path auto-engages
    axis = np.linspace(200, 1000, n_ch)
    coords = np.array([(x, y) for y in range(1, 31) for x in range(1, 51)])
    ints = []
    for _ in range(n_px):
        v = np.zeros(n_ch); idx = rng.randint(0, n_ch, 50); v[idx] = np.abs(rng.rand(50)) * 100
        ints.append(v)
    peaks = np.sort(axis[rng.randint(0, n_ch, 20)])

    class ParStore(MemoryStore):
        supports_parallel = True    # open_reader/read_with fall back to get() — safe for RAM

    def build(store):
        d = MSIDataset(store)
        d.prime(); d.build_features(peaks); d.build_mz_cube()
        return d

    ser = build(MemoryStore(coords, [axis] * n_px, ints))
    par = build(ParStore(coords, [axis] * n_px, ints))

    assert np.allclose(ser._mean[1], par._mean[1])
    for k in ("tic", "rms", "median"):
        assert np.allclose(ser._pix[k], par._pix[k]), k
    assert np.allclose(ser._feat.matrix, par._feat.matrix)
    assert np.allclose(ser._cube[1].toarray(), par._cube[1].toarray())


def test_parallel_streaming_concurrent_passes_isolated():
    """Two parallel streaming passes on the SAME store at once (e.g. the auto cube build
    racing 'Find peaks') must not share or close each other's per-thread readers."""
    import threading
    from smile_msi.msi import MemoryStore

    rng = np.random.RandomState(9)
    n_px, n_ch = 1200, 300
    axis = np.linspace(200, 1000, n_ch)
    coords = np.array([(x, y) for y in range(1, 25) for x in range(1, 51)])
    ints = [np.abs(rng.rand(n_ch)) for _ in range(n_px)]

    class ParStore(MemoryStore):
        supports_parallel = True

    ds = MSIDataset(ParStore(coords, [axis] * n_px, ints))
    ref = {i: v.copy() for i, _, v in ds._read_many(range(n_px), parallel=False)}
    errors = []

    def one_pass():
        try:
            got = {i: v for i, _, v in ds._read_many(range(n_px), parallel=True)}
            assert len(got) == n_px and all(np.array_equal(got[i], ref[i]) for i in got)
        except Exception as e:  # noqa: BLE001
            errors.append(repr(e))

    threads = [threading.Thread(target=one_pass) for _ in range(4)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert not errors, errors


def test_parse_cache_roundtrip_and_staleness(tmp_path):
    """The parse-offset sidecar lets a reopen skip the XML parse and rebuild the store from
    cached byte offsets — the spectra it reads must be byte-identical to a fresh parse (a
    wrong precision/offset would reproduce the 'buffer size' error). A changed file (mtime/
    size) must invalidate the cache and force a fresh parse."""
    import os
    from smile_msi import session

    p = str(tmp_path / "pc.imzML")
    demo.write_synthetic_imzml(p, width=14, height=11, seed=7)

    fresh = MSIDataset.from_imzml(p, lazy=True)          # first open: parses + writes sidecar
    assert fresh.store._p is not None
    assert os.path.exists(session.parse_cache_path(p))

    cached = MSIDataset.from_imzml(p, lazy=True)         # second open: rebuilt from the sidecar
    assert cached.store._p is None                       # no live parser → came from cache
    assert cached.n_pixels == fresh.n_pixels
    assert np.array_equal(cached.coordinates, fresh.coordinates)
    # every spectrum reads identically through the cached portable reader (serial + parallel)
    for i in range(fresh.n_pixels):
        m0, v0 = fresh.store.get(i)
        m1, v1 = cached.store.get(i)
        assert np.array_equal(m0, m1) and np.array_equal(v0, v1), i
    cached.prime()
    assert cached.mean_spectrum()[1].max() > 0

    # touch the .ibd → staleness key changes → cache rejected, fresh parse taken
    ibd = os.path.splitext(p)[0] + ".ibd"
    st = os.stat(ibd)
    os.utime(ibd, ns=(st.st_atime_ns, st.st_mtime_ns + 1_000_000_000))
    assert session.load_parse_cache(p) is None
    reparsed = MSIDataset.from_imzml(p, lazy=True)
    assert reparsed.store._p is not None                 # parsed fresh, not from the stale cache


def test_subsample_load(tmp_path):
    p = tmp_path / "s.imzML"
    demo.write_synthetic_imzml(str(p), width=16, height=12, seed=2)   # 192 px
    full = MSIDataset.from_imzml(str(p))
    half = MSIDataset.from_imzml(str(p), stride=2)
    assert half.n_pixels < full.n_pixels
    half.prime()
    assert half.mean_spectrum()[1].max() > 0


def test_ingest_csv_wide(tmp_path):
    import pandas as pd
    from smile_msi import ingest
    rows = []
    for y in range(4):
        for x in range(5):
            rows.append({"x": x + 1, "y": y + 1, "281.2486": x * 1.0, "327.2330": y * 1.0})
    p = tmp_path / "wide.csv"
    pd.DataFrame(rows).to_csv(p, index=False)
    ds = ingest.read_imaging_csv(str(p))
    assert ds.n_pixels == 20 and ds.width == 5 and ds.height == 4
    img = ds.ion_image(281.2486, tol_ppm=50)
    assert img[0].max() > 0          # intensity increases with x


def test_ingest_csv_long(tmp_path):
    import pandas as pd
    from smile_msi import ingest
    rows = []
    for y in range(3):
        for x in range(3):
            for mz, inten in [(281.25, 10.0 + x), (327.23, 5.0 + y)]:
                rows.append({"x": x + 1, "y": y + 1, "mz": mz, "intensity": inten})
    p = tmp_path / "long.csv"
    pd.DataFrame(rows).to_csv(p, index=False)
    ds = ingest.read_imaging_csv(str(p))
    assert ds.n_pixels == 9
    ds.prime()
    assert ds.mean_spectrum()[1].max() > 0


def test_to_imzml_roundtrip(ds, tmp_path):
    from smile_msi import ingest
    out = tmp_path / "conv.imzML"
    ingest.to_imzml(ds, str(out))
    back = MSIDataset.from_imzml(str(out))
    assert back.n_pixels == ds.n_pixels
    back.prime()
    assert back.mean_spectrum()[1].max() > 0


def test_imzml_roundtrip(tmp_path):
    p = tmp_path / "s.imzML"
    demo.write_synthetic_imzml(str(p), width=12, height=10, seed=1)
    d = MSIDataset.from_imzml(str(p), lazy=True)
    assert d.n_pixels == 120
    d.prime()
    assert d.mean_spectrum()[1].max() > 0
    assert d.polarity == "negative"


# ----- segmentation / multivariate ---------------------------------------- #
def test_segment_and_auto(ds, peaks):
    seg = spatial.segment(ds, peaks, n_clusters=4)
    assert seg.n_clusters == 4
    assert seg.labels.shape == (ds.n_pixels,)
    auto = spatial.auto_segment(ds, peaks, k_range=range(2, 6))
    assert 2 <= auto.n_clusters <= 5


def test_region_scope_mask_segments_and_components(ds, peaks):
    """A region mask (the flow's 'Restrict to:' scope) confines segmentation + PCA/NMF
    to the region's pixels: out-of-region pixels are unassigned (-1) / blank (NaN), while
    the score vectors stay full-length so the linked views don't break."""
    rows, cols = ds._pixel_rows_cols()
    mask = cols < (ds.width // 2)                       # left half = the "nerve"
    assert mask.any() and not mask.all()

    for seg in (spatial.segment(ds, peaks, n_clusters=3, mask=mask),
                spatial.auto_segment(ds, peaks, k_range=range(2, 5), mask=mask),
                multivariate.spatial_segment(ds, peaks, n_clusters=3, mask=mask)):
        assert seg.labels.shape == (ds.n_pixels,)
        assert (seg.labels[~mask] == -1).all()          # out-of-region = unassigned
        assert (seg.labels[mask] >= 0).all()            # every in-region pixel clustered
        assert np.isnan(seg.label_image[rows[~mask], cols[~mask]]).all()
        assert np.isfinite(seg.label_image[rows[mask], cols[mask]]).all()

    for comp in (multivariate.pca_images(ds, peaks, n_components=3, mask=mask),
                 multivariate.nmf_images(ds, peaks, n_components=3, mask=mask)):
        assert comp.scores.shape[0] == ds.n_pixels      # full-length (no view breakage)
        img = comp.images[0]
        assert np.isnan(img[rows[~mask], cols[~mask]]).all()
        assert np.isfinite(img[rows[mask], cols[mask]]).all()

    # mask=None is unchanged: every pixel keeps a real label
    assert (spatial.segment(ds, peaks, n_clusters=3).labels >= 0).all()


def test_hierarchy_cut_monotone_and_nested(ds, peaks):
    """One agglomerative tree, cut at any detail level (the Detail-slider engine).
    Cutting finer must yield more (or equal) segments, every pixel labelled, and the
    partition strictly refines the coarser one (each fine segment sits inside one
    coarse segment)."""
    hier = spatial.hierarchy(ds, peaks, n_micro=40)
    assert hier.micro_labels.shape == (ds.n_pixels,)
    assert hier.max_clusters == 40 and hier.linkage.shape == (39, 4)

    coarse = spatial.cut(hier, 3)
    fine = spatial.cut(hier, 7)
    assert coarse.shape == fine.shape == (ds.n_pixels,)
    assert set(np.unique(coarse)) == set(range(coarse.max() + 1))   # contiguous 0..k-1
    assert coarse.max() + 1 <= 3 and fine.max() + 1 <= 7
    assert fine.max() >= coarse.max()                               # finer = more segments
    # refinement: no fine segment straddles two coarse segments
    for f in np.unique(fine):
        assert len(np.unique(coarse[fine == f])) == 1
    # size-ordered labels: cluster 0 is the largest
    assert (np.bincount(coarse) == np.sort(np.bincount(coarse))[::-1]).all()


def test_hierarchy_cut_is_cheap_and_clamped(ds, peaks):
    """Cutting reuses the prebuilt tree (no re-embed/re-cluster) and clamps the
    requested detail to the tree's range."""
    hier = spatial.hierarchy(ds, peaks, n_micro=30)
    assert spatial.cut(hier, 1).max() == 0                          # one segment
    over = spatial.cut(hier, 999)                                   # clamped to max_clusters
    assert over.max() + 1 <= hier.max_clusters
    seg = spatial.segmentation_at(ds, hier, 4, with_silhouette=True)
    assert isinstance(seg, spatial.Segmentation)
    assert seg.label_image.shape == (ds.height, ds.width)
    assert seg.n_clusters <= 4 and np.isfinite(seg.silhouette)


def test_joint_segmentation_shared_ids_across_samples(peaks):
    """Joint segmentation builds ONE tree over two slides on a shared feature axis, so a
    cut assigns the same cluster ids to both — and each sample gets its own label image on
    its own grid. Two differently-sized synthetic slides exercise the per-sample split."""
    a = demo.make_synthetic(width=24, height=18, seed=1); a.prime()
    b = demo.make_synthetic(width=20, height=22, seed=2); b.prime()
    jh = spatial.joint_hierarchy([a, b], peaks, n_micro=40, names=["A", "B"])
    assert isinstance(jh, spatial.JointHierarchy)
    assert jh.n_samples == 2 and jh.sample_names == ["A", "B"]
    # pooled pixel count == sum of both slides; the tree spans them jointly
    assert jh.sample_sizes == [a.n_pixels, b.n_pixels]
    assert jh.hier.micro_labels.shape == (a.n_pixels + b.n_pixels,)

    segs = spatial.joint_segmentation_at([a, b], jh, 5, with_silhouette=True)
    assert len(segs) == 2
    sa, sb = segs
    assert sa.label_image.shape == (a.height, a.width)
    assert sb.label_image.shape == (b.height, b.width)
    assert sa.labels.shape == (a.n_pixels,) and sb.labels.shape == (b.n_pixels,)
    # SAME cluster vocabulary in both slides (shared ids), capped at the requested detail
    assert sa.n_clusters == sb.n_clusters <= 5
    assert set(np.unique(sa.labels)) <= set(range(sa.n_clusters))
    assert set(np.unique(sb.labels)) <= set(range(sb.n_clusters))
    assert np.isfinite(sa.silhouette)


def test_joint_segmentation_region_masks_scatter_back(peaks):
    """A per-sample mask segments only those pixels; labels scatter back onto the full
    grid with -1 off-mask (mirrors the single-slide region-scoped path)."""
    a = demo.make_synthetic(width=20, height=16, seed=3); a.prime()
    b = demo.make_synthetic(width=20, height=16, seed=4); b.prime()
    mask_a = np.zeros(a.n_pixels, dtype=bool); mask_a[: a.n_pixels // 2] = True
    jh = spatial.joint_hierarchy([a, b], peaks, n_micro=30, masks=[mask_a, None])
    assert jh.sample_sizes == [int(mask_a.sum()), b.n_pixels]
    sa, sb = spatial.joint_segmentation_at([a, b], jh, 4)
    assert (sa.labels[~mask_a] == -1).all()              # off-mask stays unlabelled
    assert (sa.labels[mask_a] >= 0).all()
    assert (sb.labels >= 0).all()                        # whole-slide sample fully labelled


def test_joint_hierarchy_release_cubes_frees_dense_keeps_geometry(peaks):
    """release_cubes frees each slide's dense cube the moment its feature block is extracted
    (the joint-seg OOM fix) — bounding resident cubes to ~1 — yet the slides' geometry
    survives so joint_segmentation_at still renders every per-slide label map."""
    a = demo.make_synthetic(width=20, height=16, seed=5); assert a.to_ram(); a.prime()
    b = demo.make_synthetic(width=18, height=14, seed=6); assert b.to_ram(); b.prime()
    jh = spatial.joint_hierarchy([a, b], peaks, n_micro=30, names=["A", "B"],
                                 release_cubes=True)
    # dense cubes freed, but n_pixels + to_image still work (coordinates outlive the matrix)
    assert getattr(a.store, "matrix", None) is None and getattr(b.store, "matrix", None) is None
    assert a.n_pixels == 20 * 16 and b.n_pixels == 18 * 14
    segs = spatial.joint_segmentation_at([a, b], jh, 4)
    assert segs[0].label_image.shape == (a.height, a.width)
    assert segs[1].label_image.shape == (b.height, b.width)
    assert segs[0].labels.shape == (a.n_pixels,)


def test_spatial_segment(ds, peaks):
    seg = multivariate.spatial_segment(ds, peaks, n_clusters=4, spatial_sigma=1.0)
    assert seg.n_clusters == 4 and seg.label_image.shape == (ds.height, ds.width)


def test_spatial_segment_auto_k(ds, peaks):
    seg = multivariate.spatial_segment(ds, peaks, n_clusters=None, k_range=range(2, 6))
    assert 2 <= seg.n_clusters <= 5


def test_spatial_feature_finding(ds, peaks):
    sa = spatial.spatial_autocorrelation(ds, peaks)
    assert {"mz", "morans_i"} <= set(sa.columns)
    assert (sa["morans_i"] <= 1.01).all()
    # the structured demo lipids should have clearly positive autocorrelation
    assert sa["morans_i"].max() > 0.3
    kept = spatial.select_spatial_features(ds, peaks, min_morans=0.1)
    assert 0 < len(kept) <= len(peaks)


def test_feature_frequency_in_unit_range(ds, peaks):
    freq = spatial.feature_frequency(ds, peaks)
    assert freq.shape == (len(peaks),)
    assert ((freq >= 0) & (freq <= 1)).all()
    # the footprint metric must report a *present* ion as present: the demo lipids
    # tile most of the tissue, so the strongest peaks have a substantial footprint.
    # (Regression guard for the old median+MAD floor, which scored reproducible ions
    # near zero and made the frequency gate cull real features.)
    assert freq.max() > 0.2


def test_feature_frequency_footprint_vs_spike():
    """A widely-present ion reports a large footprint; a single-pixel spike a tiny one."""
    H = W = 16
    axis = np.arange(299.0, 322.0, 0.01)

    def gauss(mz, amp):
        sig = mz / 30000.0 / 2.355
        a, b = np.searchsorted(axis, mz - 6 * sig), np.searchsorted(axis, mz + 6 * sig)
        out = np.zeros_like(axis)
        out[a:b] = amp * np.exp(-0.5 * ((axis[a:b] - mz) / sig) ** 2)
        return out

    coords, mzs, ints = [], [], []
    for r in range(H):
        for c in range(W):
            coords.append((c + 1, r + 1))
            spec = gauss(300.0, 1.0)                       # present in every pixel
            if (r, c) == (5, 5):
                spec += gauss(310.0, 5.0)                  # bright in exactly one pixel
            ints.append(spec.astype(np.float32))
            mzs.append(axis)
    d = MSIDataset.from_arrays(coords, mzs, ints, polarity="negative", spec_mode="profile")
    d.prime()
    freq = spatial.feature_frequency(d, [300.0, 310.0])
    assert freq[0] > 0.9                                   # ubiquitous ion → big footprint
    assert freq[1] < 0.05                                  # single-pixel spike → vanishing


def test_peak_fwhm_ppm_positive_for_real_peak(ds, peaks):
    axis, spec = ds.max_spectrum()
    w = spatial.peak_fwhm_ppm(axis, spec, peaks[0])
    assert 0 < w <= 200.0                              # bracketed, capped at the fallback


def test_find_spatial_features_pipeline(ds):
    res = spatial.find_spatial_features(ds, snr=3.0, min_frequency=0.0, min_morans=0.0)
    # funnel is monotone non-increasing and the survivors carry the new metadata
    assert res.n_candidates >= res.n_after_frequency >= res.n_after_morans == len(res.peaks)
    assert res.peaks, "demo lipids should survive a permissive run"
    assert res.params["projection"] == "mean"             # mean candidates by default
    p = res.peaks[0]
    assert {"mz", "intensity", "snr", "frequency", "morans_i", "width_ppm"} <= set(p)
    assert 0 <= p["frequency"] <= 1
    # the spatial gate genuinely prunes: a strict Moran's I cut keeps fewer features
    strict = spatial.find_spatial_features(ds, snr=3.0, min_frequency=0.0, min_morans=0.6)
    assert len(strict.peaks) <= len(res.peaks)
    assert all(q["morans_i"] >= 0.6 for q in strict.peaks)
    # survivors are ordered by descending spatial structure
    mis = [q["morans_i"] for q in res.peaks]
    assert mis == sorted(mis, reverse=True)


def test_find_spatial_features_projection_modes(ds):
    mean = spatial.find_spatial_features(ds, min_frequency=0.0, min_morans=0.0, projection="mean")
    both = spatial.find_spatial_features(ds, min_frequency=0.0, min_morans=0.0, projection="both")
    # 'both' unions mean + skyline candidates, so it can only add to the mean set
    assert both.n_candidates >= mean.n_candidates
    with pytest.raises(ValueError):
        spatial.find_spatial_features(ds, projection="bogus")


def test_find_spatial_features_denoise_keeps_signal_drops_noise():
    """The decisive property: mean candidates + Moran's denoise keep a spatially
    structured ion while dropping a spatially-random noise ion and a single-pixel
    spike — so the finder beats a plain mean threshold instead of being worse."""
    rng = np.random.default_rng(3)
    H = W = 24
    axis = np.arange(295.0, 326.0, 0.01)

    def gauss(mz, amp):
        sig = mz / 30000.0 / 2.355
        a, b = np.searchsorted(axis, mz - 6 * sig), np.searchsorted(axis, mz + 6 * sig)
        out = np.zeros_like(axis)
        if amp > 0:
            out[a:b] = amp * np.exp(-0.5 * ((axis[a:b] - mz) / sig) ** 2)
        return out

    spike_px = {(7, 7), (12, 15)}
    coords, mzs, ints = [], [], []
    for r in range(H):
        for c in range(W):
            coords.append((c + 1, r + 1))
            spec = np.zeros_like(axis)
            spec += gauss(300.0, 1.0 if c < W // 2 else 0.0)   # structured: bright left half
            spec += gauss(310.0, 0.5 * rng.random())           # i.i.d. random everywhere → noise
            if (r, c) in spike_px:
                spec += gauss(320.0, 6.0)                       # single-pixel spike
            spec += rng.random(len(axis)) * 0.01               # i.i.d. baseline
            ints.append(spec.astype(np.float32))
            mzs.append(axis)
    d = MSIDataset.from_arrays(coords, mzs, ints, polarity="negative", spec_mode="profile")
    d.prime()

    res = spatial.find_spatial_features(d, snr=3.0, min_rel_intensity=0.002,
                                        min_frequency=0.01, min_morans=0.1)
    got = np.array([p["mz"] for p in res.peaks])

    def found(mz):
        return got.size > 0 and bool(np.min(np.abs(got - mz)) <= mz * 60e-6)

    assert found(300.0), "structured ion must survive the denoise"
    assert not found(310.0), "spatially-random noise ion must be dropped (Moran's gate)"
    assert not found(320.0), "single-pixel spike must be dropped (frequency gate)"


def test_find_spatial_features_scoped_to_region(ds):
    # detection restricted to a pixel mask still returns whole-slide-extractable features
    mask = np.zeros(ds.n_pixels, dtype=bool)
    mask[: ds.n_pixels // 2] = True
    res = spatial.find_spatial_features(ds, min_frequency=0.0, min_morans=0.0, mask=mask)
    assert isinstance(res.peaks, list)
    assert res.n_candidates >= len(res.peaks)


def test_pca_nmf_components(ds, peaks):
    pca = multivariate.pca_images(ds, peaks, n_components=3)
    assert len(pca.images) == 3 and pca.explained_variance.sum() <= 1.0 + 1e-6
    nmf = multivariate.nmf_images(ds, peaks, n_components=3)
    assert len(nmf.images) == 3 and len(nmf.top_peaks(0, 3)) == 3


def test_embedding(ds, peaks):
    emb = multivariate.embedding(ds, peaks, method="tsne")
    assert emb.coords.shape == (ds.n_pixels, 2)
    assert emb.rgba.shape == (ds.height, ds.width, 4)
    assert emb.index.shape == (ds.n_pixels,)
    assert emb.point_rgb().shape == (ds.n_pixels, 3)
    # a subsampled embedding keeps coords and index aligned to the chosen pixels
    sub = multivariate.embedding(ds, peaks, method="tsne", sample=ds.n_pixels // 2)
    assert sub.coords.shape[0] == sub.index.shape[0] <= ds.n_pixels // 2


def test_pooled_embedding_pools_samples(peaks):
    """pooled_embedding stacks pixels from several samples on a shared axis and runs ONE
    embedding: every point carries a sample id + group, the per-sample cap is honoured, and
    the result lives in a single coordinate frame (cross-sample, unlike embedding())."""
    a = demo.make_synthetic(width=20, height=16, seed=1); a.prime()
    b = demo.make_synthetic(width=18, height=14, seed=2); b.prime()
    samples = [type("S", (), {"name": "ctrl-A", "group": "control"})(),
               type("S", (), {"name": "synk-B", "group": "synkinetic"})()]
    dss = {"ctrl-A": a, "synk-B": b}
    cap = 30
    emb = multivariate.pooled_embedding(
        samples, peaks, loader=lambda s: (dss[s.name], None),
        method="tsne", per_sample_cap=cap)
    assert emb.method == "TSNE"
    assert emb.sample_names == ["ctrl-A", "synk-B"]
    # each slide is capped, so the pooled cloud is exactly the two capped blocks
    assert emb.counts == {"ctrl-A": cap, "synk-B": cap}
    n = emb.coords.shape[0]
    assert n == 2 * cap
    assert emb.sample_id.shape == (n,) and emb.group.shape == (n,)
    assert set(emb.sample_id.tolist()) == {0, 1}
    assert set(emb.group[emb.sample_id == 0]) == {"control"}
    assert set(emb.group[emb.sample_id == 1]) == {"synkinetic"}


def test_pooled_embedding_honours_region_index_and_skips(peaks):
    """A region sample embeds only its pixel subset; a loader returning None skips the
    sample; an empty roster raises rather than silently embedding nothing."""
    a = demo.make_synthetic(width=20, height=16, seed=3); a.prime()
    keep = np.arange(0, a.n_pixels, 3)                      # a strided 'region'
    samples = [type("S", (), {"name": "region", "group": "g"})(),
               type("S", (), {"name": "missing", "group": "g"})()]

    def loader(s):
        return (a, keep) if s.name == "region" else None    # second sample is dropped

    emb = multivariate.pooled_embedding(samples, peaks, loader=loader, method="tsne",
                                        per_sample_cap=None)
    assert emb.sample_names == ["region"]                   # the None sample is skipped
    assert emb.coords.shape[0] == keep.size
    with pytest.raises(ValueError):
        multivariate.pooled_embedding([], peaks, loader=loader, method="tsne")


def test_pooled_embedding_keep_geometry_maps_points_to_pixels(peaks):
    """keep_geometry retains each point's display-grid (row, col) + the slide shape, so a
    cluster found in the embedding can be painted back onto tissue. Coordinates must match the
    dataset's own pixel→(row, col) accessor for the kept rows."""
    a = demo.make_synthetic(width=20, height=16, seed=7); a.prime()
    keep = np.arange(0, a.n_pixels, 2)                       # a strided subset
    samples = [type("S", (), {"name": "slideA", "group": "g"})()]
    emb = multivariate.pooled_embedding(
        samples, peaks, loader=lambda s: (a, keep), method="tsne",
        per_sample_cap=None, keep_geometry=True)
    assert emb.shapes == {"slideA": (a.height, a.width)}
    assert emb.px_row is not None and emb.px_row.shape[0] == keep.size
    rr, cc = a._pixel_rows_cols()
    assert np.array_equal(emb.px_row, np.asarray(rr, int)[keep])
    assert np.array_equal(emb.px_col, np.asarray(cc, int)[keep])
    # default (no keep_geometry) leaves the geometry fields unset
    emb2 = multivariate.pooled_embedding(samples, peaks, loader=lambda s: (a, keep),
                                         method="tsne", per_sample_cap=None)
    assert emb2.px_row is None and emb2.shapes is None


def test_pooled_embedding_errors_distinguish_load_vs_no_signal(peaks):
    """The "nothing to embed" failure names the real cause: no slide loaded (missing files)
    vs. slides loaded but every feature block was all-zero — a feature axis foreign to the
    slides. The two need opposite fixes, so the messages must differ."""
    a = demo.make_synthetic(width=20, height=16, seed=11); a.prime()
    samples = [type("S", (), {"name": "A", "group": "g"})(),
               type("S", (), {"name": "B", "group": "g"})()]

    # All loaders return None → nothing loaded.
    with pytest.raises(ValueError, match="could be loaded"):
        multivariate.pooled_embedding(samples, peaks, loader=lambda s: None, method="tsne")

    # Slides load fine, but the target axis is way off the data → all-zero extraction.
    foreign = [10000.0, 20000.0]
    with pytest.raises(ValueError, match="no target m/z fell within"):
        multivariate.pooled_embedding(samples, foreign, loader=lambda s: (a, None),
                                      method="tsne", per_sample_cap=None)


def test_cluster_points_splits_separated_blobs():
    """cluster_points groups the embedding plane: k-means returns exactly k clusters that
    recover well-separated blobs; HDBSCAN falls back to k-means when unavailable; degenerate
    sizes are handled."""
    rng = np.random.default_rng(0)
    coords = np.vstack([rng.normal((-6, -6), 0.3, (150, 2)),
                        rng.normal((6, 6), 0.3, (150, 2))])
    lab = multivariate.cluster_points(coords, method="kmeans", k=2, random_state=0)
    assert lab.shape == (300,)
    assert len(set(lab.tolist())) == 2
    # the two blobs land in different clusters (each blob is internally consistent)
    assert len(set(lab[:150].tolist())) == 1 and len(set(lab[150:].tolist())) == 1
    # HDBSCAN path: returns one label per point (real clusters if installed, else k-means)
    hd = multivariate.cluster_points(coords, method="hdbscan", k=2, random_state=0)
    assert hd.shape == (300,)
    # edge cases
    assert multivariate.cluster_points(np.zeros((0, 2))).shape == (0,)
    assert multivariate.cluster_points(np.zeros((1, 2))).tolist() == [0]


def test_features_for_rows_matches_full_subset_and_release():
    """features_for_rows extracts only the chosen pixels (no whole-slide build/cache) and is
    numerically identical to feature_matrix(norm)[rows] — the rank-2 extraction fix. And
    MSIDataset.release() frees the dense cube + caches (the cohort OOM eviction primitive)."""
    ds = demo.make_synthetic(width=18, height=14, seed=11)
    pk = [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)][:8]
    assert len(pk) >= 2
    rows = np.array([0, 5, 17, 99, 100, 200], dtype=int)
    rows = rows[rows < ds.n_pixels]
    assert ds.to_ram()                                   # exercise the dense optimization path
    for norm in ("none", "tic"):
        full = multivariate_feature_matrix(ds, pk, norm)
        sub = ds.features_for_rows(pk, rows, tol_ppm=20.0, norm=norm)
        assert sub.shape == (rows.size, len(pk))
        assert np.allclose(sub, full[rows], rtol=1e-5, atol=1e-6)

    ds.release()
    assert ds._feat is None
    assert getattr(ds.store, "matrix", None) is None     # dense cube freed deterministically


def multivariate_feature_matrix(ds, peaks, norm):
    from smile_msi import spatial
    return spatial.feature_matrix(ds, peaks, tol_ppm=20.0, norm=norm)


def test_features_for_rows_streams_only_requested_pixels():
    """Without a dense cube, features_for_rows STREAMS only the requested pixels (no whole-slide
    pass) — the region-mean speed fix that keeps a single big slide from loading entirely.
    Raw (norm='none') extraction is bit-for-bit the full path subset; tic stays finite."""
    ds = demo.make_synthetic(width=18, height=14, seed=5)        # lazy — NOT to_ram'd → streaming
    assert ds._dense() is None                                   # confirm the streaming branch runs
    pk = [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)][:6]
    assert len(pk) >= 2
    rows = np.array([1, 4, 33, 77, 150], dtype=int)
    rows = rows[rows < ds.n_pixels]
    sub_none = ds.features_for_rows(pk, rows, tol_ppm=20.0, norm="none")
    assert sub_none.shape == (rows.size, len(pk))
    assert np.allclose(sub_none, multivariate_feature_matrix(ds, pk, "none")[rows],
                       rtol=1e-5, atol=1e-6)                     # raw extraction is exact
    sub_tic = ds.features_for_rows(pk, rows, tol_ppm=20.0, norm="tic")
    assert sub_tic.shape == (rows.size, len(pk)) and np.isfinite(sub_tic).all()


def test_coords_rgb_wheel_reaches_green():
    """The 2-D embedding colour wheel must reach the full hue spectrum — including a
    clearly green colour — so green regions aren't hidden in the scatter the way the old
    x→R/y→G/complement→B mapping hid them (only visible on the tissue map)."""
    ang = np.linspace(0, 2 * np.pi, 24, endpoint=False)
    coords = np.column_stack([np.cos(ang), np.sin(ang)])
    rgb = multivariate._coords_rgb(coords)
    assert rgb.shape == (24, 3) and rgb.min() >= 0 and rgb.max() <= 1
    # at least one point is dominantly green (G clearly above both R and B)
    green = rgb[:, 1] - np.maximum(rgb[:, 0], rgb[:, 2])
    assert green.max() > 0.4
    assert multivariate._coords_rgb(np.zeros((0, 2))).shape == (0, 3)   # empty-input guard


def test_pooled_region_embedding_one_point_per_region(peaks):
    """pooled_region_embedding embeds ONE point per named region per sample (the FTU-
    instance-level view): each point carries its region label, sample id, group and the
    pixel count behind its mean; an empty roster raises."""
    a = demo.make_synthetic(width=20, height=16, seed=11); a.prime()
    b = demo.make_synthetic(width=18, height=14, seed=12); b.prime()
    dss = {"A": a, "B": b}
    samples = [type("S", (), {"name": "A", "group": "facial"})(),
               type("S", (), {"name": "B", "group": "facial"})()]

    def regions_of(s):
        ds = dss[s.name]
        half = ds.n_pixels // 2
        return [("fascicle", np.arange(0, half)),
                ("perineurium", np.arange(half, ds.n_pixels))]

    emb = multivariate.pooled_region_embedding(
        samples, peaks, loader=lambda s: (dss[s.name], None), regions_of=regions_of,
        unit="region", method="tsne")
    assert emb.unit == "region"
    assert emb.coords.shape == (4, 2)                       # 2 samples x 2 regions
    assert sorted(set(emb.region.tolist())) == ["fascicle", "perineurium"]
    assert set(emb.sample_id.tolist()) == {0, 1}
    assert emb.sizes.shape == (4,) and (emb.sizes > 0).all()
    assert emb.counts == {"A": 2, "B": 2}
    with pytest.raises(ValueError):
        multivariate.pooled_region_embedding([], peaks, loader=lambda s: None,
                                             regions_of=regions_of)


def test_pooled_region_embedding_instance_unit_splits_blobs(peaks):
    """unit='instance' splits a region into spatially connected components, so two
    disjoint blobs of one region become two points (one nerve fascicle each)."""
    a = demo.make_synthetic(width=24, height=12, seed=13); a.prime()
    rows, cols = a._pixel_rows_cols()
    # one region = two spatially separated vertical strips (a gap between them)
    left = np.flatnonzero(cols <= 2)
    right = np.flatnonzero(cols >= a.width - 3)
    region_idx = np.concatenate([left, right])
    sample = type("S", (), {"name": "A", "group": "g"})()

    def regions_of(_s):
        return [("fascicles", region_idx)]

    region = multivariate.pooled_region_embedding(
        [sample, type("S", (), {"name": "B", "group": "g"})()], peaks,
        loader=lambda s: (a, None), regions_of=regions_of, unit="region", method="tsne")
    assert region.coords.shape[0] == 2                      # one mean per sample

    inst = multivariate.pooled_region_embedding(
        [sample], peaks, loader=lambda s: (a, None), regions_of=regions_of,
        unit="instance", method="tsne")
    assert inst.coords.shape[0] == 2                        # two connected blobs -> two points
    assert (inst.region == "fascicles").all()


# ----- statistics ---------------------------------------------------------- #
def test_auc_mwu_matches_scipy():
    # validate the vectorized rank-AUC + Mann-Whitney against scipy's reference.
    # AUC is the exact rank effect size; the p-value uses method='auto' (asymptotic at
    # these large n) WITH continuity correction — matching scipy's defaults.
    from scipy.stats import mannwhitneyu
    from smile_msi.spatial import _auc_mwu
    rng = np.random.default_rng(0)
    XA = rng.normal(0.0, 1.0, (40, 4))
    XB = rng.normal(0.6, 1.0, (55, 4))
    auc, p = _auc_mwu(XA, XB)
    for j in range(4):
        U, pp = mannwhitneyu(XB[:, j], XA[:, j], alternative="two-sided",
                             use_continuity=True, method="auto")
        assert abs(auc[j] - U / (40 * 55)) < 1e-9       # AUC == U/(na*nb) exactly
        assert abs(p[j] - pp) < 1e-9                     # p matches scipy (auto + continuity)


def test_auc_mwu_small_n_is_exact_not_anticonservative():
    # the headline HIGH-4/11 fix: at small, tie-free sample sizes the p-value must be the
    # exact MWU null (≈0.10 for 3-vs-3 perfect separation), not the old asymptotic 0.0495
    # that crossed 0.05 and produced false positives in the sample-summarized paths.
    import numpy as np
    from smile_msi.spatial import _auc_mwu
    XA = np.array([[1.0], [2.0], [3.0]])
    XB = np.array([[4.0], [5.0], [6.0]])               # B strictly greater than A
    auc, p = _auc_mwu(XA, XB)
    assert auc[0] == 1.0                                # perfect separation, high in B
    assert p[0] >= 0.1 - 1e-9                           # exact null, not the 0.0495 approx


def test_roc_curve_matches_rank_auc():
    # the empirical ROC's area must equal the rank-based AUC (>0.5 ⇒ higher in B), and the
    # curve must run monotonically from (0,0) to (1,1).
    from smile_msi.spatial import roc_curve, _auc_mwu
    rng = np.random.default_rng(1)
    va = rng.normal(0.0, 1.0, 80)
    vb = rng.normal(0.7, 1.0, 95)
    fpr, tpr, auc = roc_curve(va, vb)
    auc_ref, _ = _auc_mwu(va.reshape(-1, 1), vb.reshape(-1, 1))
    assert abs(auc - float(auc_ref[0])) < 1e-9 and auc > 0.5
    assert fpr[0] == 0 and tpr[0] == 0
    assert abs(fpr[-1] - 1) < 1e-12 and abs(tpr[-1] - 1) < 1e-12
    assert np.all(np.diff(fpr) >= -1e-12) and np.all(np.diff(tpr) >= -1e-12)
    # swapping the groups mirrors the curve below the diagonal (AUC < 0.5)
    _, _, auc_swap = roc_curve(vb, va)
    assert abs(auc_swap - (1.0 - auc)) < 1e-9
    # empty / degenerate group → the chance diagonal
    f0, _t0, a0 = roc_curve(np.array([]), vb)
    assert a0 == 0.5 and list(f0) == [0.0, 1.0]


def test_bh_fdr_matches_scipy():
    from scipy.stats import false_discovery_control
    from smile_msi.spatial import _bh_fdr
    ps = np.array([0.001, 0.008, 0.02, 0.04, 0.2, 0.5, 0.011, 0.3])
    assert np.allclose(_bh_fdr(ps), false_discovery_control(ps, method="bh"), atol=1e-9)


def test_roi_comparison_exact_auc(ds, peaks):
    seg = spatial.segment(ds, peaks, n_clusters=3)
    res = spatial.roi_comparison(ds, seg.mask(0), seg.mask(1), peaks)
    assert (res["AUC"].between(0, 1)).all()
    assert (res["q_value"] >= res["p_value"] - 1e-9).all()  # FDR >= raw p


def test_discriminating_and_multigroup(ds, peaks):
    seg = spatial.segment(ds, peaks, n_clusters=3)
    disc = spatial.discriminating_features(ds, seg.labels, peaks, top_n=5)
    assert set(disc.keys()) <= {0, 1, 2} and all(len(d) <= 5 for d in disc.values())
    mg = spatial.multigroup_features(ds, seg.labels, peaks)
    assert {"mz", "H", "p_value", "q_value", "top_region"} <= set(mg.columns)


def test_discriminating_presence_gate():
    """A sparse, spiky ion can post a high enriched-AUC from rank alone yet have no
    real signal across the region — it must be flagged ``present=False`` (so it sorts
    below true markers and is excluded from a region's 'markers' list), while a marker
    that's actually on across the region is ``present=True``."""
    axis = np.arange(295.0, 336.0, 0.01)

    def gauss(mz, amp):
        sig = mz / 30000.0 / 2.355
        a, b = np.searchsorted(axis, mz - 6 * sig), np.searchsorted(axis, mz + 6 * sig)
        out = np.zeros_like(axis)
        if amp > 0:
            out[a:b] = amp * np.exp(-0.5 * ((axis[a:b] - mz) / sig) ** 2)
        return out

    N = 40
    g0 = set(range(0, 20))                      # "region" A pixels; the rest are B
    spike_px = {2, 5, 8}                         # only 3 of A's 20 pixels carry m/z 330
    coords, mzs, ints = [], [], []
    for c in range(N):
        spec = gauss(320.0, 1.0)                 # 320: ubiquitous (not discriminating)
        if c in g0:
            spec += gauss(300.0, 1.0)            # 300: real A marker — on across all of A
            if c in spike_px:
                spec += gauss(330.0, 1.0)        # 330: sparse spikes in A only
        else:
            spec += gauss(310.0, 1.0)            # 310: real B marker
        coords.append((c + 1, 1))
        mzs.append(axis)
        ints.append(spec.astype(np.float32))
    d = MSIDataset.from_arrays(coords, mzs, ints, polarity="negative", spec_mode="profile")
    d.prime()

    labels = np.array([0 if c in g0 else 1 for c in range(N)])
    pk = [300.0, 310.0, 320.0, 330.0]
    disc = spatial.discriminating_features(d, labels, pk, top_n=4, present_frac=0.5)
    a = disc[0].set_index(disc[0]["mz"].round(1))

    assert "present" in disc[0].columns
    # both 300 and 330 read as enriched in A by AUC...
    assert a.loc[300.0, "AUC"] > 0.5 and a.loc[330.0, "AUC"] > 0.5
    # ...but only the real marker is detected across the region
    assert bool(a.loc[300.0, "present"]) is True
    assert bool(a.loc[330.0, "present"]) is False
    # present ions rank first, so the genuine marker outranks the sparse spike
    order = list(disc[0]["mz"].round(1))
    assert order.index(300.0) < order.index(330.0)


def test_aggregate_by_class():
    # pure roll-up: summing each class's ion columns, unannotated ('') ions dropped
    X = np.array([[1.0, 2.0, 10.0, 5.0],
                  [3.0, 4.0, 20.0, 7.0]])
    Xc, names, n_ions = spatial._aggregate_by_class(X, ["PE", "PE", "PC", ""])
    assert names == ["PC", "PE"] and n_ions == [1, 2]
    assert np.allclose(Xc[:, 0], [10.0, 20.0])           # PC = the single PC ion
    assert np.allclose(Xc[:, 1], [3.0, 7.0])             # PE = ion0 + ion1
    with pytest.raises(ValueError):                       # no annotated lipids at all
        spatial._aggregate_by_class(X, ["", "", "", ""])
    with pytest.raises(ValueError):                       # classes not parallel to columns
        spatial._aggregate_by_class(X, ["PE", "PC"])


def test_class_comparison_rolls_up_to_classes(ds, peaks):
    from smile_msi.match import Annotator
    classes = Annotator(mode="negative", ppm_tol=10).classes_for(peaks)
    n_annot = sum(1 for c in classes if c)
    uniq = sorted(set(c for c in classes if c))
    assert len(uniq) >= 2                                  # the demo spans several classes
    seg = spatial.segment(ds, peaks, n_clusters=3)
    ma, mb = seg.mask(0), seg.mask(1)
    res = spatial.class_comparison(ds, ma, mb, peaks, classes)
    assert list(res["class"]) == uniq                      # one row per annotated class
    assert {"class", "n_ions", "mean_A", "mean_B", "AUC", "log2_fc",
            "p_value", "q_value"} <= set(res.columns)
    assert int(res["n_ions"].sum()) == n_annot             # every annotated ion counted once
    assert res["AUC"].between(0, 1).all()
    assert (res["q_value"] >= res["p_value"] - 1e-9).all()
    assert res.attrs["kind"] == "class"
    # AUC is symmetric: swapping A/B mirrors it about 0.5 (higher-in-B becomes higher-in-A)
    rev = spatial.class_comparison(ds, mb, ma, peaks, classes)
    assert np.allclose(res["AUC"].to_numpy(), 1.0 - rev["AUC"].to_numpy())


def test_class_composition_percent_per_region(ds, peaks):
    from smile_msi.match import Annotator
    classes = Annotator(mode="negative", ppm_tol=10).classes_for(peaks)
    uniq = sorted(set(c for c in classes if c))
    seg = spatial.segment(ds, peaks, n_clusters=3)
    comp = spatial.class_composition(ds, [seg.mask(0), seg.mask(1)], peaks, classes,
                                     names=["WM", "GM"])
    assert list(comp.index) == uniq and list(comp.columns) == ["WM", "GM"]
    assert np.allclose(comp.sum(axis=0).to_numpy(), 100.0)  # each region sums to 100%
    assert (comp.to_numpy() >= 0).all()
    assert comp.attrs["regions"] == ["WM", "GM"]
    assert set(comp.attrs["n_ions"]) == set(uniq)
    assert list(comp.attrs["intensity"].index) == uniq


def test_roi_comparison_parametric_methods(ds, peaks):
    from scipy.stats import ttest_ind
    seg = spatial.segment(ds, peaks, n_clusters=3)
    ma, mb = seg.mask(0), seg.mask(1)
    base = spatial.roi_comparison(ds, ma, mb, peaks)            # default = MWU
    assert base.attrs["test"].startswith("exact rank-based")
    X = spatial.feature_matrix(ds, peaks, norm="tic")
    for method, equal_var, label in [("welch", False, "Welch's t-test (unequal variance)"),
                                     ("student", True, "Student's t-test (pooled variance)")]:
        res = spatial.roi_comparison(ds, ma, mb, peaks, method=method)
        assert res.attrs["test"] == label
        # AUC is unchanged (still rank-based) but p now matches scipy's t-test
        assert np.allclose(res["AUC"], base["AUC"])
        _, p = ttest_ind(X[ma], X[mb], axis=0, equal_var=equal_var)
        assert np.allclose(res["p_value"], np.where(np.isfinite(p), p, 1.0))
        assert (res["q_value"] >= res["p_value"] - 1e-9).all()
    with pytest.raises(ValueError):
        spatial.roi_comparison(ds, ma, mb, peaks, method="bogus")


def test_multigroup_anova_matches_scipy(ds, peaks):
    from scipy.stats import f_oneway
    seg = spatial.segment(ds, peaks, n_clusters=3)
    mg = spatial.multigroup_features(ds, seg.labels, peaks, method="anova")
    assert {"mz", "F", "p_value", "q_value", "top_region"} <= set(mg.columns)
    assert mg.attrs["test"] == "one-way ANOVA F"
    X = spatial.feature_matrix(ds, peaks, norm="tic")
    groups = [X[seg.labels == c] for c in np.unique(seg.labels)]
    row = mg.iloc[0]                                            # top (lowest p) ion
    j = int(np.argmin(np.abs(np.asarray(peaks, float) - row["mz"])))
    F, p = f_oneway(*[g[:, j] for g in groups])
    assert abs(row["F"] - F) < 1e-6 and abs(row["p_value"] - p) < 1e-9
    with pytest.raises(ValueError):
        spatial.multigroup_features(ds, seg.labels, peaks, method="bogus")


def test_multigroup_top_region_uses_group_names(ds, peaks):
    # When the caller passes group names aligned to the label ids, top_region reads the
    # readable name (e.g. 'Cortex'), not the raw integer cluster id the CSV used to show.
    seg = spatial.segment(ds, peaks, n_clusters=3)
    names = ["Cortex", "Medulla", "Pelvis"]
    mg = spatial.multigroup_features(ds, seg.labels, peaks, names=names)
    assert set(mg["top_region"]).issubset(set(names))
    assert mg["top_region"].map(lambda v: isinstance(v, str)).all()
    # names=None keeps the raw integer ids (back-compat for callers that don't have names)
    mg_int = spatial.multigroup_features(ds, seg.labels, peaks)
    assert set(mg_int["top_region"]).issubset(set(np.unique(seg.labels).tolist()))


def test_registry_per_group_analyses_carry_group_names(ds, peaks):
    # The multi-group and marker (shrunken-centroid) steps must surface the assigned group
    # names in their result tables, not raw integer region ids — the reported CSV bug.
    from smile_msi import registry
    seg = spatial.segment(ds, peaks, n_clusters=3)
    names = ["Cortex", "Medulla", "Pelvis"]
    inp = {"labels": seg.labels, "mzs": peaks, "names": names}
    p = {"tol_ppm": 20.0, "norm": "tic", "method": "kruskal", "shrink": 2.0, "top_n": 10}

    mg = registry._run_multigroup(ds, inp, p)
    assert set(mg["top_region"]).issubset(set(names))

    sh = registry._run_shrunken(ds, inp, p)
    tbl = registry._per_cluster_table(sh)
    assert set(tbl["region"]).issubset(set(names))

    # segmentation fallback (no assigned groups → names None) stays labelled 'region <n>'
    sh_seg = registry._run_shrunken(ds, {**inp, "names": None}, p)
    tbl_seg = registry._per_cluster_table(sh_seg)
    assert all(str(r).startswith("region ") for r in tbl_seg["region"])


def test_per_region_analysis_ignores_unassigned_pixels(ds, peaks):
    # ROI-style labelling that doesn't tile the slide: -1 pixels must be dropped, so the
    # result has exactly the two real groups (not a spurious third "-1" group).
    seg = spatial.segment(ds, peaks, n_clusters=3)
    labels = np.full(ds.n_pixels, -1)
    labels[seg.mask(0)] = 0
    labels[seg.mask(1)] = 1
    disc = spatial.discriminating_features(ds, labels, peaks, top_n=5)
    assert set(disc.keys()) == {0, 1}
    mg = spatial.multigroup_features(ds, labels, peaks)
    assert set(mg["top_region"]).issubset({0, 1})
    # group sizes recorded for the small-sample / power guard
    for cl, df in disc.items():
        assert df.attrs["n_in"] == int((labels == cl).sum())
        assert df.attrs["n_in"] + df.attrs["n_out"] == int((labels >= 0).sum())
    assert mg.attrs["group_sizes"] == {0: int((labels == 0).sum()), 1: int((labels == 1).sum())}


def test_region_membership_venn(ds, peaks):
    seg = spatial.segment(ds, peaks, n_clusters=3)
    masks = [seg.mask(0), seg.mask(1), seg.mask(2)]
    res = spatial.region_membership(ds, masks, peaks, names=["A", "B", "C"],
                                    min_prevalence=0.5, detect_quantile=0.75)
    assert res["names"] == ["A", "B", "C"] and len(res["n_pixels"]) == 3
    # every compartment's key is a subset of the region indices; counts match the m/z lists
    for c in res["compartments"]:
        assert set(c["key"]) <= {0, 1, 2} and c["count"] == len(c["mzs"])
        assert c["exclusive"] == (len(c["key"]) == 1)
    # no ion is assigned to two compartments at once
    assigned = [mz for c in res["compartments"] for mz in c["mzs"]]
    assert len(assigned) == len(set(assigned))
    # optional significance gate must not increase the assigned set
    gated = spatial.region_membership(ds, masks, peaks, max_q=0.05)
    assert sum(c["count"] for c in gated["compartments"]) <= len(peaks)
    assert "q" in gated["peaks"].columns


def test_colocalize_and_matrix(ds, peaks):
    co = spatial.colocalize(ds, peaks[0], peaks)
    assert abs(co[0]["score"] - 1.0) < 1e-6            # target self-correlates ~1
    M, mpk = spatial.coloc_matrix(ds, peaks)
    assert M.shape == (len(peaks), len(peaks))
    assert np.allclose(np.diag(M), 1.0, atol=1e-6)


def test_coloc_modules(ds, peaks):
    mods = spatial.coloc_modules(ds, peaks, n_modules=4)
    assert mods.matrix.shape == (len(peaks), len(peaks))
    assert len(mods.labels) == len(peaks)
    members = mods.members()
    assert 1 <= len(members) <= 4
    assert sum(len(v) for v in members.values()) == len(peaks)   # every ion assigned once


# ----- imaging helpers ----------------------------------------------------- #
def test_quantile_clip_and_rgb(ds, peaks):
    img = ds.ion_image(peaks[0])
    clipped = imaging.quantile_clip(img, high=95)
    assert clipped.max() <= img.max()
    rgb = imaging.rgb_overlay(ds, peaks[0], peaks[1], peaks[2])
    assert rgb.shape == (ds.height, ds.width, 3) and rgb.dtype == np.uint8


def test_export_ion_figure(ds, peaks, tmp_path):
    out = tmp_path / "ion.png"
    imaging.export_ion_figure(ds.ion_image(peaks[0]), str(out), title="m/z test",
                              pixel_size_um=20.0, scale_bar_um=200.0)
    assert out.exists() and out.stat().st_size > 0


def test_colormap_and_optical_overlay(ds, peaks):
    img = ds.ion_image(peaks[0])
    rgba = imaging.apply_colormap(img, cmap="viridis")        # exercises matplotlib colormap API
    assert rgba.shape == (ds.height, ds.width, 4) and rgba.dtype == np.uint8
    optical = np.random.randint(0, 255, (30, 40, 3), dtype=np.uint8)
    blend = imaging.overlay_optical(rgba, optical, alpha=0.5)
    assert blend.shape == (ds.height, ds.width, 3) and blend.dtype == np.uint8


# ----- preprocessing ------------------------------------------------------- #
def test_preprocessing_pipeline_preserves_axis(ds):
    n_axis = len(ds.mean_spectrum()[0])
    d = demo.make_synthetic(width=12, height=10, seed=3)
    d.set_preprocessing(preprocess.build_pipeline({"baseline": 20, "smooth": {"savgol": 9}}))
    d.prime()
    axis, spec = d.mean_spectrum()
    assert spec.max() > 0                              # signal survives baseline+smooth
    assert (spec >= 0).all()


def test_recalibration_runs(ds):
    d = demo.make_synthetic(width=10, height=8, seed=4)
    from smile_msi.demo import _demo_peaks
    refs = [m for m, _w, _n in _demo_peaks()[:3]]
    d.set_preprocessing([preprocess.recalibrate(refs, tol_ppm=300)])
    d.prime()
    assert d.mean_spectrum()[1].max() > 0


# ----- annotation ---------------------------------------------------------- #
def test_isotope_consistency(ds, peaks):
    iso = isotopes.isotope_consistency(ds, peaks[0])
    assert iso["m1_ratio"] > 0 and iso["consistent"]   # demo has isotopes


def test_provenance(ds, tmp_path):
    from smile_msi import provenance
    p = provenance.Provenance(title="t", started="2026-06-14T00:00:00+00:00")
    p.set_dataset(ds, "negative")
    p.step("peak_picking", snr=3.0, norm="tic", now="2026-06-14T00:00:01+00:00")
    p.step("segmentation", spatial=True, auto=True, n_clusters=4, now="2026-06-14T00:00:02+00:00")
    d = p.to_dict()
    assert d["dataset"]["pixels"] == ds.n_pixels
    assert [s["step"] for s in d["steps"]] == ["peak_picking", "segmentation"]
    assert "numpy" in d["environment"]
    md = p.to_markdown(str(tmp_path / "methods.md"))
    assert "Provenance" in md and "Methods (draft)" in md
    para = p.methods_paragraph()
    assert "segmented" in para and "Peaks were detected" in para
    p.to_json(str(tmp_path / "prov.json"))
    assert (tmp_path / "prov.json").exists() and (tmp_path / "methods.md").exists()


def test_file_fingerprint(tmp_path):
    from smile_msi import provenance
    f = tmp_path / "x.bin"
    f.write_bytes(b"hello world")
    fp = provenance.file_fingerprint(str(f))
    assert fp["bytes"] == 11 and len(fp["sha256"]) == 64 and fp["method"] == "sha256"


def test_report_with_provenance_sheet(ds, tmp_path):
    from smile_msi import provenance, pipeline
    pk = ds.pick_peaks(snr=3, min_rel_intensity=0.01)
    mzs = [p["mz"] for p in pk]
    seg = spatial.segment(ds, mzs, n_clusters=3)
    res = spatial.roi_comparison(ds, seg.mask(0), seg.mask(1), mzs)
    res = pipeline.annotate_df(res, "mz", mode="negative", ppm=10)
    prov = provenance.Provenance(started="2026-06-14T00:00:00+00:00").set_dataset(ds, "negative")
    prov.step("segmentation", n_clusters=3)
    out = tmp_path / "r.xlsx"
    pipeline.build_report(res, str(out), a_label="A", b_label="B", provenance=prov)
    from openpyxl import load_workbook
    wb = load_workbook(str(out))
    assert "Provenance & methods" in wb.sheetnames


def test_feature_list_and_fdr(ds):
    pk = ds.pick_peaks(snr=3, min_rel_intensity=0.01)
    df = annotate.build_feature_list(ds, pk, mode="negative")
    assert (df["lipid"] != "").sum() >= 8
    assert {"confidence", "isotope_ok", "n_adducts", "pubmed_url", "isotopologue"} <= set(df.columns)
    fdr = annotate.estimate_fdr([p["mz"] for p in pk], mode="negative")
    assert 0.0 <= fdr["fdr"] <= 1.0 and fdr["n_peaks"] == len(pk)


def test_feature_list_confidence_columns(ds):
    """build_feature_list emits a 0–100 confidence_score + a confidence_why breakdown."""
    pk = ds.pick_peaks(snr=3, min_rel_intensity=0.01)
    df = annotate.build_feature_list(ds, pk, mode="negative")
    assert {"confidence_score", "confidence_why"} <= set(df.columns)
    idd = df[df["lipid"] != ""]
    scores = [int(s) for s in idd["confidence_score"] if s != ""]
    assert scores and all(0 <= s <= 100 for s in scores)
    assert (idd["confidence_why"].str.len() > 0).all()


def test_confidence_detail_blend():
    # tight ppm + isotope + sibling adduct + clear winner → High
    hi = isotopes.confidence_detail(0.5, True, 2, ppm_tol=10.0, score_gap=2.0)
    assert hi["label"] == "High" and 0 <= hi["score"] <= 100 and hi["reasons"]
    # loose ppm, no isotope/adduct support, isobaric tie → Low, and clearly lower
    lo = isotopes.confidence_detail(9.5, False, 1, ppm_tol=10.0, score_gap=0.1)
    assert lo["label"] == "Low" and lo["score"] < hi["score"]
    assert isotopes.confidence(0.5, True, 2, 10.0, 2.0) == "High"        # back-compat wrapper


def test_fdr_randomized_decoy_reliability(ds):
    pk = [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)]
    fdr = annotate.estimate_fdr(pk, mode="negative")
    assert "reliable" in fdr and 0.0 <= fdr["fdr"] <= 1.0
    # deterministic: identical input reproduces the decoy hit-rate exactly
    assert annotate.estimate_fdr(pk, mode="negative")["decoy_rate"] == fdr["decoy_rate"]
    # a 2-peak list is flagged as too short to trust
    assert annotate.estimate_fdr(pk[:2], mode="negative")["reliable"] is False


def test_msms_ganglioside_neu5ac(tmp_path):
    """The sialic-acid B-ion (290.0881) confirms a ganglioside class."""
    from smile_msi import msms
    p = tmp_path / "g.mgf"
    p.write_text("BEGIN IONS\nPEPMASS=772.0\n290.0881 9000\n264.2686 3000\nEND IONS\n")
    spec = msms.parse_mgf(str(p))[0]
    assert "GM1" in msms.matched_classes(spec, "negative")
    assert msms.confirm_class("GM1", spec, "negative") == "confirmed"


def test_msms_confirmation(tmp_path):
    from smile_msi import msms
    mgf = ("BEGIN IONS\nPEPMASS=888.62\n96.9601 9000\n281.2486 4000\n255.2330 2000\nEND IONS\n")
    p = tmp_path / "s.mgf"
    p.write_text(mgf)
    sp = msms.parse_mgf(str(p))[0]
    assert sp.precursor == 888.62 and len(sp.mz) == 3
    assert msms.confirm_class("Sulfatide", sp, "negative") == "confirmed"
    assert set(msms.acyl_chains(sp)) >= {"16:0", "18:1"}
    # ranked by carboxylate intensity: 18:1 (4000) before 16:0 (2000)
    assert msms.acyl_chains(sp, max_chains=2)[0] == "18:1"
    # serine neutral loss confirms PS
    ps = tmp_path / "ps.mgf"
    ps.write_text("BEGIN IONS\nPEPMASS=788.5447\n701.5127 5000\nEND IONS\n")
    assert msms.confirm_class("PS", msms.parse_mgf(str(ps))[0], "negative") == "confirmed"
    # expanded rules: lyso headgroup, hexose loss, sphingoid base
    lpi = msms.MS2Spectrum(599.32, np.array([241.0119, 152.9958]), np.array([5e3, 2e3]))
    assert msms.confirm_class("LPI", lpi, "negative") == "confirmed"
    hc = msms.MS2Spectrum(888.62, np.array([888.62 - 162.0528]), np.array([4e3]))
    assert msms.confirm_class("HexCer", hc, "negative") == "confirmed"
    sm = msms.MS2Spectrum(703.57, np.array([184.0733, 264.2686]), np.array([9e3, 3e3]))
    assert msms.confirm_class("SM", sm, "positive") == "confirmed"
    # classify an unknown spectrum -> PI ranks top (3 diagnostic ions present)
    unknown = msms.MS2Spectrum(885.55, np.array([241.0119, 223.0013, 152.9958]),
                               np.array([5e3, 3e3, 2e3]))
    ranked = msms.classify(unknown, "negative")
    assert ranked and ranked[0][0] == "PI"


def test_session_roundtrip(ds, tmp_path):
    from smile_msi import session
    pk = ds.pick_peaks(snr=3, min_rel_intensity=0.01)
    seg = spatial.segment(ds, [p["mz"] for p in pk], n_clusters=3)
    mask_a = seg.mask(0)
    s = session.build_session(
        source=ds.source, settings={"mode": "negative", "ppm": 50, "norm": "tic"},
        peaks=pk, active_mz=pk[0]["mz"], labels=seg.labels, n_clusters=seg.n_clusters,
        region_a=mask_a, region_b=seg.mask(1),
        feature_lists={"Panel A": [{"mz": 700.5, "lipid": "PC", "note": "n"}]},
        feature_scopes={"All slide": pk}, active_feature_scope="All slide",
        flist_name="All slide",
        report_items=[{"type": "ion", "title": "PC 34:1", "caption": "", "mz": 700.5,
                       "ppm": 10.0, "window": [0.0, 99.0], "cmap": "viridis"},
                      {"type": "note", "title": "Intro", "text": "Curated report."}])
    p = tmp_path / "s.json"
    session.save_session(str(p), s)
    back = session.load_session(str(p))
    assert back["source"] == ds.source
    assert len(back["peaks"]) == len(pk)
    assert back["segmentation"]["n_clusters"] == 3
    m = session.mask_from_indices(back["regions"]["A"], ds.n_pixels)
    assert np.array_equal(m, mask_a)
    # per-sample analysis store (feature lists + working scopes) round-trips
    assert list(back["feature_lists"]) == ["Panel A"]
    assert back["feature_lists"]["Panel A"][0]["mz"] == 700.5
    assert len(back["feature_scopes"]["All slide"]) == len(pk)
    assert back["active_feature_scope"] == "All slide"
    assert back["flist_name"] == "All slide"
    # curated PDF-report contents (Report tab) survive the round-trip in order
    assert [it["type"] for it in back["report_items"]] == ["ion", "note"]
    assert back["report_items"][0]["mz"] == 700.5
    # a session written without report_items still loads (defaulted to an empty list)
    assert session.load_session(str(p)).get("report_items") is not None


def test_session_named_region_sample_roundtrip(tmp_path):
    """A named region's owning slide (its 'sample' tag) persists; a legacy region saved
    without one round-trips as '' so it loads as 'the current slide' — no migration."""
    from smile_msi import session
    regions = [{"name": "Tumor", "color": "#ff0000", "segments": {0, 1},
                "visible": True, "group": "A", "sample": "/d/slideX.imzML"},
               {"name": "Legacy", "color": "#00ff00", "segments": {2}, "visible": True}]
    s = session.build_session(source="/d/slideX.imzML", settings={"mode": "negative"},
                              peaks=[{"mz": 700.0, "intensity": 1.0}], named_regions=regions)
    p = tmp_path / "s.json"
    session.save_session(str(p), s)
    nr = {r["name"]: r for r in session.load_session(str(p))["named_regions"]}
    assert nr["Tumor"]["sample"] == "/d/slideX.imzML"     # owning slide persists
    assert nr["Legacy"]["sample"] == ""                   # missing tag → "" (current slide)


# --------------------------------------------------------------------------- #
# full-pipeline integration
# --------------------------------------------------------------------------- #
def test_end_to_end_workflow(tmp_path):
    from smile_msi import pipeline, provenance
    # 1. write + lazy-load imzML
    p = tmp_path / "e2e.imzML"
    demo.write_synthetic_imzml(str(p), width=24, height=20, seed=9)
    d = MSIDataset.from_imzml(str(p))
    d.prime()
    # 2. peaks + features + fast cube
    pk = [x["mz"] for x in d.pick_peaks(snr=3, min_rel_intensity=0.01)]
    assert len(pk) >= 10
    d.ensure_features(pk)
    d.build_mz_cube()
    assert d.ion_image(pk[0] + 0.6).shape == (d.height, d.width)          # arbitrary m/z via cube
    # 3. segmentation (auto + spatially-aware)
    seg = spatial.auto_segment(d, pk)
    assert seg.n_clusters >= 2
    assert multivariate.spatial_segment(d, pk, n_clusters=None).n_clusters >= 2
    # 4. exact stats -> annotate -> Excel report with provenance
    res = pipeline.annotate_df(spatial.roi_comparison(d, seg.mask(0), seg.mask(1), pk),
                               "mz", mode="negative", ppm=10)
    assert (res["best_lipid"] != "").sum() >= 5
    prov = provenance.Provenance(started="2026-06-14T00:00:00+00:00").set_dataset(d, "negative")
    out = tmp_path / "report.xlsx"
    pipeline.build_report(res, str(out), a_label="A", b_label="B", provenance=prov)
    assert out.exists() and out.stat().st_size > 5000
    # 5. feature list, discriminating features, modules, multigroup
    feat = annotate.build_feature_list(d, [{"mz": m} for m in pk], mode="negative")
    assert len(feat) == len(pk)
    assert spatial.discriminating_features(d, seg.labels, pk)
    assert spatial.coloc_modules(d, pk, n_modules=3).members()
    assert len(spatial.multigroup_features(d, seg.labels, pk)) == len(pk)


# --------------------------------------------------------------------------- #
# robustness / edge cases
# --------------------------------------------------------------------------- #
def test_robustness_flat_dataset():
    axis = np.linspace(200, 900, 50)
    coords = [(x + 1, y + 1) for y in range(3) for x in range(3)]
    ds0 = MSIDataset.from_arrays(coords, [axis] * 9, [np.zeros(50) for _ in coords])
    ds0.prime()
    assert ds0.pick_peaks() == []                     # no peaks in flat data
    img = ds0.ion_image(500.0)                         # all-zero ion image, no crash
    assert img.shape == (3, 3) and img.max() == 0


def test_pick_peaks_collapses_shoulder_notch():
    """A single strong, well-resolved peak with a shallow one-bin dip at its apex
    (detector quantization / mild saturation ripple) used to be picked as two adjacent
    local maxima a fraction of a mDa apart — both clearing the height/prominence gate
    since that gate is a per-window noise estimate, not a saddle-depth test. Each twin
    then annotated to the same lipid within tolerance, producing exact-duplicate rows in
    the feature list. ``pick_peaks`` must collapse them to the one real peak."""
    n = 4000
    axis = np.linspace(780.0, 790.0, n)
    rng = np.random.default_rng(7)
    i0 = 2000
    coords = [(1, 1), (2, 1), (1, 2), (2, 2)]
    ints = []
    for _ in coords:
        spec = np.abs(rng.normal(0, 2.0, size=n))
        spec[i0 - 1] = 4998.0
        spec[i0] = 4990.0          # the shallow notch
        spec[i0 + 1] = 5000.0
        ints.append(spec)
    ds = MSIDataset.from_arrays(coords, [axis] * len(coords), ints)
    ds.prime()
    pk = ds.pick_peaks(snr=3, min_rel_intensity=0.002, prominence=1.0)
    near = [p for p in pk if 784.9 < p["mz"] < 785.1]
    assert len(near) == 1, f"expected the notch to collapse to one peak, got {near}"


def test_robustness_tiny_dataset():
    axis = np.linspace(200, 900, 40)
    coords = [(1, 1), (2, 1), (1, 2), (2, 2)]
    ints = [np.exp(-0.5 * ((axis - 281.0) / 0.3) ** 2) * (i + 1) for i in range(4)]
    tiny = MSIDataset.from_arrays(coords, [axis] * 4, ints)
    tiny.prime()
    pk = tiny.pick_peaks(snr=1.0)
    assert len(pk) >= 1
    tiny.ensure_features([p["mz"] for p in pk])
    assert tiny.feature_matrix("tic").shape[0] == 4


def test_robustness_empty_fdr_and_deisotope():
    assert annotate.estimate_fdr([])["fdr"] == 1.0
    mono, is_iso = isotopes.deisotope([])
    assert mono == [] and len(is_iso) == 0


def test_robustness_empty_inputs_raise(ds, peaks):
    with pytest.raises(ValueError):
        spatial.segment(ds, [], n_clusters=3)
    seg = spatial.segment(ds, peaks, n_clusters=3)
    with pytest.raises(ValueError):                    # empty region
        spatial.roi_comparison(ds, np.zeros(ds.n_pixels, bool), seg.mask(0), peaks)


def test_robustness_single_peak_analytics(ds, peaks):
    sa = spatial.spatial_autocorrelation(ds, [peaks[0]])
    assert len(sa) == 1
    mods = spatial.coloc_modules(ds, [peaks[0]])
    assert len(mods.members()) == 1
    assert imaging.quantile_clip(np.zeros((5, 6)), high=99).max() == 0


def test_wider_database():
    from smile_msi.lipiddb import build_database
    idx = {l.name: l for l in build_database()}
    classes = {l.lipid_class for l in build_database()}
    assert {"GM1", "GM3", "GD1", "CE", "ST", "CAR", "Metab"} <= classes
    # masses verified against known monoisotopic values (< 5 ppm)
    for name, ref in [("GM3 36:1", 1180.7438), ("GM1 36:1", 1545.8758),
                      ("CE 18:1", 650.6006), ("Cholesterol", 386.3549),
                      ("CAR 16:0", 399.3349), ("N-acetylaspartate (NAA)", 175.0481),
                      ("Taurine", 125.0147), ("Glutathione (GSH)", 307.0838)]:
        assert abs(idx[name].neutral_mass - ref) / ref * 1e6 < 5
    assert "FA 22:6;O1" in idx          # oxylipins (oxidized fatty acids)


def test_deisotope(ds):
    # the demo adds M+1/M+2 to every base lipid -> deisotoping must drop ~2/3 of peaks
    pk = ds.pick_peaks(snr=3, min_rel_intensity=0.01)
    mono, is_iso = isotopes.deisotope(pk)
    assert is_iso.sum() > 0
    assert len(mono) < len(pk)
    # monoisotopic peaks should be the more intense ones
    mono_mz = {float(p["mz"]) for p in mono}
    from smile_msi.demo import _demo_peaks
    for mz, _w, _name in _demo_peaks():
        assert min(abs(m - mz) for m in mono_mz) < mz * 30 / 1e6  # base lipids kept


def test_session_unique_feature_list_name():
    from smile_msi import session
    assert session.unique_feature_list_name([], "Lesion only") == "Lesion only"
    assert session.unique_feature_list_name(["Lesion only"], "Lesion only") == "Lesion only (2)"
    assert session.unique_feature_list_name(
        ["Lesion only", "Lesion only (2)"], "Lesion only") == "Lesion only (3)"
    assert session.unique_feature_list_name([], "") == "Feature list"


def test_dataset_fingerprint_distinguishes_slides(ds):
    """Two different-sized slides must fingerprint differently; the same slide stable."""
    from smile_msi import library
    fp = library.dataset_fingerprint(ds)
    assert fp and fp == library.dataset_fingerprint(ds)        # deterministic
    other = demo.make_synthetic(width=40, height=30, seed=7)
    assert library.dataset_fingerprint(other) != fp           # different slide → different id


def test_session_fingerprint_mismatch_detection(ds):
    """A restored session detects when it's opened on a different slide than it was
    saved on — by fingerprint (preferred) or, for legacy sessions, by pixel count."""
    from smile_msi import session, library
    fp = library.dataset_fingerprint(ds)
    assert not session.fingerprint_mismatch(fp, fp)             # same slide
    assert session.fingerprint_mismatch(fp, fp + "x")           # different fingerprint
    # legacy session without a fingerprint → fall back to pixel count
    assert not session.fingerprint_mismatch(None, None, ds.n_pixels, ds.n_pixels)
    assert session.fingerprint_mismatch(None, None, ds.n_pixels, ds.n_pixels + 1)
    # unknown on one side → never cry wolf
    assert not session.fingerprint_mismatch("", fp)


def test_session_records_dataset_identity(ds):
    from smile_msi import session, library
    s = session.build_session(
        source=ds.source, settings={"mode": "negative"}, peaks=[{"mz": 700.5}],
        n_pixels=ds.n_pixels, dataset_fingerprint=library.dataset_fingerprint(ds))
    assert s["n_pixels"] == ds.n_pixels
    assert s["dataset_fingerprint"] == library.dataset_fingerprint(ds)


def test_isotope_scores_discriminate(ds):
    """Real ion -> high spectral+spatial (consistent); empty m/z -> low (not)."""
    from smile_msi.match import Annotator
    pk = ds.pick_peaks(snr=3, min_rel_intensity=0.01)
    mz = pk[0]["mz"]
    best = Annotator("negative", 10.0).annotate_mz(mz)[0]
    pat = isotopes.theoretical_isotope_pattern(best.lipid.formula, best.adduct, 3)
    real = isotopes.isotope_scores(ds, pat, tol_ppm=50)
    assert real["spectral"] > 0.8 and real["consistent"]
    # a pattern anchored on an empty m/z gap should not look like a real isotope ladder
    gap = float(ds.mean_spectrum()[0].max()) - 0.37
    fake = [(gap, 1.0), (gap + 1.00336, 0.5), (gap + 2.0067, 0.12)]
    assert isotopes.isotope_scores(ds, fake, tol_ppm=50)["consistent"] is False


def test_fdr_target_decoy_qvalues(ds):
    """estimate_fdr returns per-peak q-values aligned to input + FDR-tier counts."""
    mzs = [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)]
    f = annotate.estimate_fdr(mzs, mode="negative", ds=ds, n_decoy=6)
    assert len(f["q_values"]) == len(mzs)
    assert set(f["levels"]) == {0.05, 0.10, 0.20, 0.50}
    assert f["image_based"] and f["n_decoy"] == 6
    qs = [q for q in f["q_values"] if q == q]
    assert qs and all(0.0 <= q <= 1.0 for q in qs)


def test_fdr_ranking_separates_real_from_decoy():
    """The rank-based FDR + rule-of-succession yields a low q when targets clearly
    out-score decoys, and a high q when they overlap (mass-only fallback path)."""
    # build a peak list of exact real ion m/z (targets score perfectly on mass)
    from smile_msi.lipiddb import build_database
    from smile_msi.masses import ion_mz
    db = build_database()
    real_mzs = [ion_mz(l.neutral_mass, "[M-H]-") for l in db
                if l.lipid_class in ("PE", "PS", "PI", "PG")][:40]
    f = annotate.estimate_fdr(real_mzs, mode="negative", n_decoy=20)
    # exact-mass targets (q ~ low) should mostly pass a 50% screen
    assert f["levels"][0.50] >= f["levels"][0.05]
    assert 0.0 <= f["fdr"] <= 1.0


def test_roi_comparison_sample_summarized(ds):
    """Sample-summarized testing avoids pseudoreplication: across-replicate p is far
    less extreme than the per-pixel p, and the result is tagged unit='sample'."""
    pk = [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)]
    seg = spatial.segment(ds, pk, n_clusters=3)
    a, b = seg.mask(0), seg.mask(1)
    pp = spatial.roi_comparison(ds, a, b, pk)
    samples = np.full(ds.n_pixels, -1, int)
    for grp, base in [(a, 0), (b, 3)]:
        for k, part in enumerate(np.array_split(np.flatnonzero(grp), 3)):
            samples[part] = base + k
    ss = spatial.roi_comparison(ds, a, b, pk, samples=samples)
    assert pp.attrs["unit"] == "pixel" and ss.attrs["unit"] == "sample"
    assert ss.attrs["n_a"] == 3 and ss.attrs["n_b"] == 3
    assert ss["p_value"].min() >= pp["p_value"].min()
    # one replicate per group -> underpowered warning
    one = np.where(a, 0, np.where(b, 1, -1))
    w = spatial.roi_comparison(ds, a, b, pk, samples=one)
    assert "warning" in w.attrs and w.attrs["n_a"] == 1


def test_colocalization_measures(ds):
    """Co-localization measures all return finite, bounded scores; an ion is
    perfectly co-localized with itself."""
    pk = [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)]
    # cor/cosine/moc/dice are 1.0 for a perfectly co-localized (identical) image;
    # Manders m1/m2 give the fraction of signal above the median footprint (<1 by design)
    for m, lo in (("pearson", 0.99), ("cosine", 0.99), ("moc", 0.99),
                  ("dice", 0.99), ("m1", 0.5), ("m2", 0.5)):
        res = spatial.colocalize(ds, pk[0], pk, method=m)
        top = next(d for d in res if abs(d["mz"] - pk[0]) < 1e-6)
        assert top["score"] > lo, (m, top["score"])
        assert all(-1.0001 <= d["score"] <= 1.0001 for d in res if d["score"] == d["score"])


def test_manders_dice_handle_binary_and_uniform_ions():
    """Regression for the median-of-positives footprint collapsing under strict '>'.
    A uniform/binary ion's positive pixels all equal the median, so '>' gave an empty
    footprint → M2/Dice = 0/NaN even for perfectly co-localized ions; '>=' fixes it."""
    from smile_msi.spatial import _similarity
    # ion a fully nested in ion b (a present ⊂ b present), both uniform where present
    a = np.array([0.0, 0.0, 5.0, 5.0, 5.0, 5.0])
    b = np.array([0.0, 0.0, 9.0, 9.0, 9.0, 9.0])
    assert _similarity(a, b, "m1") == pytest.approx(1.0)        # all of a's signal where b present
    assert _similarity(b, a, "m2") == pytest.approx(1.0)        # all of a's signal where a present
    assert _similarity(a, b, "dice") == pytest.approx(1.0)      # identical footprints
    # a binary ion is perfectly co-localized with itself
    binary = np.array([0.0, 1.0, 0.0, 1.0, 1.0, 0.0])
    for m in ("m1", "m2", "dice"):
        assert _similarity(binary, binary, m) == pytest.approx(1.0), m
    # an all-zero image yields an empty footprint, not "everything"
    zero = np.zeros(6)
    assert not np.isfinite(_similarity(zero, b, "m2")) or _similarity(zero, b, "m2") == 0.0


def test_spatial_weights_sa_vs_sasa(ds):
    """Both spatial-weight kernels segment; adaptive (bilateral/SASA) is edge-aware."""
    pk = [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)]
    sg = multivariate.spatial_segment(ds, pk, n_clusters=4, weights="gaussian")
    sa = multivariate.spatial_segment(ds, pk, n_clusters=4, weights="adaptive")
    assert len(sg.labels) == ds.n_pixels and len(sa.labels) == ds.n_pixels
    assert sg.n_clusters == 4 and sa.n_clusters == 4


def test_shrunken_centroids_sparsity(ds):
    """Soft-threshold shrinkage selects discriminating ions; more shrink -> fewer."""
    pk = [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)]
    seg = spatial.segment(ds, pk, n_clusters=4)
    lo = spatial.shrunken_centroids(ds, seg.labels, pk, shrink=1.0)
    hi = spatial.shrunken_centroids(ds, seg.labels, pk, shrink=5.0)
    n_lo = sum(d.attrs["n_selected"] for d in lo.values())
    n_hi = sum(d.attrs["n_selected"] for d in hi.values())
    assert n_lo > 0 and n_hi <= n_lo
    for d in lo.values():                       # retained ions are non-zero, sorted
        assert (d["shrunken"] != 0).all()


def test_embedding_reports_method_and_warns(monkeypatch):
    """UMAP requested but unavailable -> t-SNE + a warning (no silent method swap)."""
    import sys
    from smile_msi import multivariate as mv
    monkeypatch.setitem(sys.modules, "umap", None)      # force the import to fail
    coords, used = mv._reduce_2d(np.random.RandomState(0).rand(60, 5), "umap", 0)
    assert used == "tsne" and coords.shape == (60, 2)


def test_spatial_dgmm_segments_ion(ds):
    """spatialDGMM splits one ion image into ordered intensity components; the hot
    component is brightest and spatially coherent."""
    pk = [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)]
    labels, means = multivariate.spatial_dgmm(ds, pk[0], k=3)
    assert len(labels) == ds.n_pixels and len(means) == 3
    assert np.all(np.diff(means) >= -1e-9)              # components ordered low->high
    img = ds.ion_vector(pk[0], norm="tic")
    assert img[labels == labels.max()].mean() > img[labels == 0].mean()


def test_plsda_classifies_and_vip(ds):
    """PLS-DA recovers a 2-cluster labelling with high train accuracy and ranks ions by VIP;
    OPLS-DA applies orthogonal filtering on binary problems."""
    pk = [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)]
    seg = spatial.segment(ds, pk, n_clusters=2)
    clf = multivariate.plsda(ds, pk, seg.labels, n_components=2)
    X = spatial.feature_matrix(ds, pk, norm="tic")
    assert float(np.mean(clf.predict(X) == seg.labels)) > 0.85
    assert clf.vip.shape == (len(pk),) and len(clf.top_peaks(3)) == 3
    oclf = multivariate.plsda(ds, pk, seg.labels, n_components=2, orthogonal=True)
    assert oclf.method == "OPLS-DA" and len(oclf._ortho) >= 1


def test_cross_validate_leakage_safe(ds):
    """Leave-one-sample-out CV groups whole samples (leakage_safe); pixel folds warn."""
    pk = [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)]
    seg = spatial.segment(ds, pk, n_clusters=2)
    samples = np.full(ds.n_pixels, -1, int)
    for cl in (0, 1):
        for k, part in enumerate(np.array_split(np.flatnonzero(seg.labels == cl), 3)):
            samples[part] = cl * 3 + k
    cv = multivariate.cross_validate(ds, pk, seg.labels, samples=samples, n_components=2)
    assert cv["leakage_safe"] and 0.0 <= cv["accuracy"] <= 1.0
    assert cv["confusion"].sum() > 0
    px = multivariate.cross_validate(ds, pk, seg.labels, n_components=2)
    assert px["leakage_safe"] is False and "warning" in px


def test_cross_validate_single_sample_falls_back(ds):
    """A single sample-group (one slide, or all but one region/sample deselected) used to
    crash LeaveOneGroupOut ('fewer than 2 groups'). It must now fall back to stratified
    pixel folds and flag the result leakage-prone instead of raising."""
    pk = [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)]
    seg = spatial.segment(ds, pk, n_clusters=3)
    samples = np.full(ds.n_pixels, 7, int)          # one and only one sample
    cv = multivariate.cross_validate(ds, pk, seg.labels, samples=samples, n_components=2)
    assert cv["leakage_safe"] is False
    assert 0.0 <= cv["accuracy"] <= 1.0
    assert "one sample" in cv["warning"]


def test_segmentation_test_sample_level(ds):
    """segmentationTest (spatialDGMM + across-sample meansTest) returns sample-unit q-values."""
    pk = [p["mz"] for p in ds.pick_peaks(snr=3, min_rel_intensity=0.01)]
    seg = spatial.segment(ds, pk, n_clusters=2)
    samples = np.full(ds.n_pixels, -1, int)
    for cl in (0, 1):
        for k, part in enumerate(np.array_split(np.flatnonzero(seg.labels == cl), 3)):
            samples[part] = cl * 3 + k
    groups = np.where(seg.labels == 0, "ctrl", "test")
    st = multivariate.segmentation_test(ds, pk[:6], groups, samples, k=3)
    assert set(["mz", "statistic", "p_value", "q_value"]) <= set(st.columns)
    assert st.attrs["unit"] == "sample" and len(st) == 6
    assert ((st["q_value"] >= 0) & (st["q_value"] <= 1)).all()


def test_provenance_records_more_methods():
    """Provenance now drafts methods + refs for multi-group, Venn, classification, and
    preprocessing/normalize steps (previously dormant), and records outputs."""
    from smile_msi import provenance
    p = provenance.Provenance(started="2026-06-15T00:00:00+00:00")
    p.step("preprocess", baseline={"method": "locmin"}, now="2026-06-15T00:00:01+00:00")
    p.step("multigroup", test="Kruskal-Wallis", now="2026-06-15T00:00:02+00:00")
    p.step("region_membership", n_compartments=4, now="2026-06-15T00:00:03+00:00")
    p.step("classification", method="PLS-DA", now="2026-06-15T00:00:04+00:00")
    p.output("/tmp/report.pdf", "data book")
    para = p.methods_paragraph()
    assert "Kruskal" in para and "Venn" in para and "PLS-DA" in para
    refs = p.references_used()
    assert any("Kruskal" in r for r in refs) and any("PLS-regression" in r for r in refs)
    assert p.to_dict()["outputs"] and p.to_dict()["outputs"][0]["file"] == "report.pdf"


# --------------------------------------------------------------------------- #
# _mask_key — compact blake2b cache key for mean_spectrum (audit plan 23, item E)
# --------------------------------------------------------------------------- #
def test_mask_key_distinguishes_masks():
    """Different masks → different keys, incl. single-pixel and same-popcount cases."""
    from smile_msi.msi import _mask_key
    base = np.zeros(40, dtype=bool)
    base[:10] = True
    one_off = base.copy(); one_off[20] = True          # +1 pixel (different popcount)
    same_count = np.zeros(40, dtype=bool); same_count[10:20] = True   # same popcount, moved
    keys = {_mask_key(base), _mask_key(one_off), _mask_key(same_count)}
    assert len(keys) == 3


def test_masked_mean_caches_and_hits_without_recompute():
    """Distinct masks → 2 cache entries with distinct results; repeating a mask is a
    cache hit served without re-reading spectra."""
    d = demo.make_synthetic(width=18, height=14, seed=21)
    d.prime()
    coords = d.coordinates
    m1 = coords[:, 0] < coords[:, 0].mean()
    m2 = ~m1
    r1 = d.mean_spectrum(m1)
    r2 = d.mean_spectrum(m2)
    assert len(d._masked_mean_cache) == 2
    assert not np.allclose(r1[1], r2[1])               # different regions differ

    calls = {"n": 0}
    orig = d._read_many
    def spy(idxs):
        calls["n"] += 1
        return orig(idxs)
    d._read_many = spy
    again = d.mean_spectrum(m1)
    assert calls["n"] == 0                              # served from cache, no read
    assert again is r1                                  # same cached object


def test_masked_mean_value_matches_fresh_dataset():
    """The key change must not alter the *value* — only the lookup."""
    d1 = demo.make_synthetic(width=16, height=12, seed=5); d1.prime()
    d2 = demo.make_synthetic(width=16, height=12, seed=5); d2.prime()
    mask = d1.coordinates[:, 1] < d1.coordinates[:, 1].mean()
    a = d1.mean_spectrum(mask)
    b = d2.mean_spectrum(mask)
    assert np.allclose(a[0], b[0]) and np.allclose(a[1], b[1])


# --------------------------------------------------------------------------- #
# Plan 18 Issue A — signed log2_fc unification (cross-function sign agreement)
# --------------------------------------------------------------------------- #
def test_fold_tau_empty_and_allzero_is_one():
    from smile_msi.spatial import _fold_tau
    assert _fold_tau(np.zeros((4, 3))) == 1.0
    assert _fold_tau(np.empty((0, 3))) == 1.0
    X = np.array([[0.0, 2.0], [0.0, 4.0], [0.0, 100.0]])
    assert _fold_tau(X) > 0                       # 5th pct of the nonzero column


def _two_region_ds():
    """3 ions over 20 pixels: ion0 high in region B, ion1 high in region A, ion2 ~absent
    in both. Region A = first 10 pixels, B = last 10."""
    axis = np.array([500.0, 600.0, 700.0])
    rng = np.random.default_rng(0)
    rows = []
    for i in range(20):
        b = i >= 10
        i0 = (50.0 if b else 2.0) + rng.uniform(0, 1)     # higher in B
        i1 = (2.0 if b else 50.0) + rng.uniform(0, 1)     # higher in A
        i2 = 0.0                                            # absent in both regions
        rows.append(np.array([i0, i1, i2]))
    coords = np.array([[x % 5 + 1, x // 5 + 1] for x in range(20)], dtype=int)
    ds = MSIDataset.from_arrays(coords, [axis] * 20, rows)
    ds.prime()
    mask_a = np.array([i < 10 for i in range(20)])
    return ds, [500.0, 600.0, 700.0], mask_a


def test_log2_fc_sign_agrees_across_functions():
    ds, peaks, mask_a = _two_region_ds()
    mask_b = ~mask_a
    roi = spatial.roi_comparison(ds, mask_a, mask_b, peaks, norm="none")
    # roi_comparison: log2_fc = log2((mB+tau)/(mA+tau)); ion0 higher in B ⇒ > 0
    by_mz = {round(float(m), 1): float(v) for m, v in zip(roi["mz"], roi["log2_fc"])}
    assert by_mz[500.0] > 0.5          # enriched in B
    assert by_mz[600.0] < -0.5         # enriched in A (depleted in B)
    assert by_mz[700.0] == 0.0         # absent in both ⇒ exactly log2(1) = 0

    # discriminating_features cluster B (label 1, "in"=B): ion0 enriched ⇒ same + sign
    labels = np.where(mask_a, 0, 1)
    disc = spatial.discriminating_features(ds, labels, peaks, norm="none", present_frac=0.0)
    dfb = disc[1]
    dmz = {round(float(m), 1): float(v) for m, v in zip(dfb["mz"], dfb["log2_fc"])}
    assert dmz[500.0] > 0              # cross-function sign agreement (the property A guarantees)
    assert dmz[600.0] < 0


# --------------------------------------------------------------------------- #
# Physical pixel size + scale-bar accuracy
# --------------------------------------------------------------------------- #
def _grid_ds(xs, ys):
    coords = np.array([(x, y) for y in ys for x in xs])
    spec = [np.array([100.0, 200.0]), np.array([1.0, 2.0])]
    return MSIDataset.from_arrays(coords, [spec[0]] * len(coords), [spec[1]] * len(coords))


def test_set_pixel_size_override_and_clear():
    ds = _grid_ds(range(5), range(4))
    assert ds.pixel_size_um is None and ds.pixel_size_y_um is None
    ds.set_pixel_size(40)
    assert ds.pixel_size_um == 40.0 and ds.pixel_size_y_um == 40.0      # y defaults to x
    ds.set_pixel_size(40, 80)
    assert ds.pixel_size_um == 40.0 and ds.pixel_size_y_um == 80.0
    ds.set_pixel_size(0)                                                # non-positive clears
    assert ds.pixel_size_um is None and ds.pixel_size_y_um is None


def test_set_pixel_size_survives_to_ram():
    ds = _grid_ds(range(5), range(4))
    ds.set_pixel_size(25)
    assert ds.to_ram()                                                  # dense, continuous mode
    assert ds.pixel_size_um == 25.0 and ds.store.pixel_size_um == 25.0


def test_coordinate_grid_step_unit_and_strided():
    assert _grid_ds(range(5), range(4)).coordinate_grid_step() == (1, 1)
    assert _grid_ds(range(0, 10, 2), range(4)).coordinate_grid_step() == (2, 1)
    # degenerate single-column grid → x step falls back to 1 (no spacing to infer)
    assert _grid_ds([3], range(4)).coordinate_grid_step() == (1, 1)


def test_pixel_size_warning_flags_nonsquare_and_strided():
    sq = _grid_ds(range(5), range(4))
    sq.set_pixel_size(40)
    assert sq.pixel_size_warning() is None                             # square + unit grid: silent

    ns = _grid_ds(range(5), range(4))
    ns.set_pixel_size(40, 80)
    assert "non-square" in ns.pixel_size_warning()

    strided = _grid_ds(range(0, 10, 2), range(4))
    strided.set_pixel_size(50)
    assert "step" in strided.pixel_size_warning()
