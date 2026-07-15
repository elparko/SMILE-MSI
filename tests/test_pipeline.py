"""Pipeline tests run on the committed synthetic sample (no real data needed)."""
import io
import os

import pandas as pd

from smile_msi import pipeline

SAMPLE = os.path.join(os.path.dirname(__file__), "..", "examples", "sample_peaks.csv")
COLS = dict(mz_col="mz", a_mean="Control_mean", a_sd="Control_sd",
            b_mean="Treatment_mean", b_sd="Treatment_sd")


def test_read_table():
    df = pipeline.read_table(SAMPLE)
    assert len(df) == 9
    assert "mz" in df.columns


def test_read_filelike_and_delimiter_sniff():
    # bytes buffer with ';' delimiter and a leading comment line
    raw = b"# exported\nmz;A;B\n100.0;1;2\n200.0;3;4\n"
    buf = io.BytesIO(raw); buf.name = "x.csv"
    df = pipeline.read_table(buf)
    assert list(df.columns) == ["mz", "A", "B"] and len(df) == 2


def test_label_from_column():
    assert pipeline.label_from_column("Control_mean") == "Control"
    assert pipeline.label_from_column("Treatment_sd") == "Treatment"
    assert pipeline.label_from_column("Group A average intensity") == "Group A"
    assert pipeline.label_from_column("mean") == ""            # nothing left
    assert pipeline.label_from_column("mean", "Group A") == "Group A"  # falls back
    assert pipeline.label_from_column(None, "Group B") == "Group B"


def test_auc_direction():
    df = pipeline.read_table(SAMPLE)
    s = pipeline.compute_auc(df, **COLS)
    auc = dict(zip(s.mz.round(4), s.AUC))
    assert auc[281.2486] < 0.05      # strongly group A (Control)
    assert auc[892.6207] > 0.90      # strongly group B (Treatment)
    assert 0.6 < auc[463.2832] < 0.9    # moderately group B (Treatment)


def test_annotation_ids_and_unidentified():
    df = pipeline.read_table(SAMPLE)
    res = pipeline.annotate_and_score(df, **COLS)
    by_mz = dict(zip(res.mz.round(4), res.best_lipid))
    assert by_mz[281.2486] == "FA 18:1"
    assert by_mz[255.2330] == "FA 16:0"
    assert by_mz[303.2330] == "FA 20:4"   # arachidonic
    assert by_mz[763.3694] == ""          # unidentified
    # every identified feature carries an MS1 id_confidence (0–100); the unidentified one is NaN
    conf = dict(zip(res.mz.round(4), res.id_confidence))
    assert 0 <= conf[281.2486] <= 100
    assert conf[763.3694] != conf[763.3694]           # NaN for the unidentified m/z


def test_annotate_df_recalibrates_shifted_comparison():
    """A comparison table over a mis-calibrated dataset recovers the shifted ions when
    recalibration is on (same gated self-calibration as the Features tab)."""
    import pandas as pd

    from smile_msi.lipiddb import build_database
    from smile_msi.masses import ion_mz
    db = build_database()
    trues = sorted({round(ion_mz(l.neutral_mass, "[M-H]-"), 5) for l in db
                    if l.lipid_class in ("PE", "PC", "PS", "PI", "FA", "Sulfatide")})[::7][:40]
    obs = [t * (1 + 4e-6) for t in trues]             # +4 ppm miscalibration
    df = pd.DataFrame({"mz": obs})
    off = pipeline.annotate_df(df, "mz", ppm=2.0, recalibrate=False)
    on = pipeline.annotate_df(df, "mz", ppm=2.0, recalibrate=True)
    n_off = int((off["best_lipid"] != "").sum())
    n_on = int((on["best_lipid"] != "").sum())
    assert n_on > n_off + 20                           # shifted ions recovered
    assert abs(on.attrs["mass_offset_ppm"] - 4.0) < 0.5
    assert off.attrs["mass_offset_ppm"] == 0.0


def test_build_report_to_buffer():
    df = pipeline.read_table(SAMPLE)
    res = pipeline.annotate_and_score(df, **COLS)
    buf = io.BytesIO()
    pipeline.build_report(res, buf, a_label="Control", b_label="Treatment")
    assert buf.getbuffer().nbytes > 5000
    from openpyxl import load_workbook
    buf.seek(0)
    wb = load_workbook(buf)
    assert wb.sheetnames[:4] == ["Overview", "Identified lipids", "Unidentified peaks",
                                 "Key & abbreviations"]
    assert wb["_data"].sheet_state == "hidden"        # chart source data is tucked away

    # four native charts live on the Overview sheet
    import zipfile
    buf.seek(0)
    chart_parts = [n for n in zipfile.ZipFile(buf).namelist()
                   if n.startswith("xl/charts/chart") and n.endswith(".xml")]
    assert len(chart_parts) == 4
